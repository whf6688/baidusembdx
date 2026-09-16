from io import BytesIO

from openpyxl import Workbook

import search_console.imports as imports_module
from search_console.imports import (
    MAX_KEYWORDS_PER_UNIT,
    filter_new_material_rows,
    parse_material_workbook,
    preview_workbook,
)


def workbook_bytes(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or ["计划", "关键词"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_preview_requires_only_campaign_and_keyword_and_skips_duplicates():
    content = workbook_bytes([
        ["计划A", "减肥"],
        ["计划A", "减肥"],
    ])
    preview = preview_workbook("material.xlsx", content)
    assert preview.row_count == 2
    assert preview.valid_count == 1
    assert preview.invalid_count == 0
    assert preview.skip_count == 1
    assert preview.errors == []
    assert preview.target_accounts == []
    assert preview.create_count == 1


def test_preview_rejects_missing_material_columns():
    preview = preview_workbook(
        "material.xlsx",
        workbook_bytes([["计划A"]], headers=["计划"]),
    )
    assert preview.invalid_count == 1
    assert "关键词" in preview.errors[0]["message"]


def test_single_pass_parser_returns_preview_and_normalized_rows(monkeypatch):
    calls = 0
    original_load_workbook = imports_module.load_workbook

    def counted_load_workbook(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(imports_module, "load_workbook", counted_load_workbook)
    content = workbook_bytes([["计划A", "减肥"], ["计划B", "控制体重"]])
    preview, rows = parse_material_workbook("material.xlsx", content)
    assert calls == 1
    assert preview.valid_count == 2
    assert rows == [
        {
            "row_number": 2,
            "campaign_name": "计划A",
            "keyword_text": "减肥",
            "keyword_utf8_encoded": "%E5%87%8F%E8%82%A5",
            "keyword_encoding_version": "utf8-rfc3986-v1",
        },
        {
            "row_number": 3,
            "campaign_name": "计划B",
            "keyword_text": "控制体重",
            "keyword_utf8_encoded": "%E6%8E%A7%E5%88%B6%E4%BD%93%E9%87%8D",
            "keyword_encoding_version": "utf8-rfc3986-v1",
        },
    ]


def test_project_material_library_appends_only_new_keywords():
    preview, rows = parse_material_workbook(
        "任意文件名.xlsx",
        workbook_bytes([
            ["计划A", "减肥"],
            ["计划A", "控制体重"],
            ["计划B", "健康饮食"],
        ]),
    )

    filtered_preview, new_rows = filter_new_material_rows(
        preview,
        rows,
        {"减肥"},
    )

    assert [(row["campaign_name"], row["keyword_text"]) for row in new_rows] == [
        ("计划A", "控制体重"),
        ("计划B", "健康饮食"),
    ]
    assert new_rows[0]["keyword_utf8_encoded"] == "%E6%8E%A7%E5%88%B6%E4%BD%93%E9%87%8D"
    assert new_rows[0]["keyword_encoding_version"] == "utf8-rfc3986-v1"
    assert filtered_preview.valid_count == 2
    assert filtered_preview.create_count == 2
    assert filtered_preview.skip_count == 1
    assert {item["campaign"] for item in filtered_preview.unit_chunks} == {"计划A", "计划B"}


def test_project_material_library_ignores_filename_and_case_for_duplicates():
    first_preview, first_rows = parse_material_workbook(
        "第一批.xlsx",
        workbook_bytes([["Plan A", "Keyword A"]]),
    )
    second_preview, second_rows = parse_material_workbook(
        "完全不同的文件名.xlsx",
        workbook_bytes([["plan a", "keyword a"]]),
    )

    first_result, first_new_rows = filter_new_material_rows(first_preview, first_rows, set())
    second_result, second_new_rows = filter_new_material_rows(
        second_preview,
        second_rows,
        {first_new_rows[0]["keyword_text"]},
    )

    assert first_result.valid_count == 1
    assert second_new_rows == []
    assert second_result.valid_count == 0
    assert second_result.skip_count == 1


def test_import_collapses_legacy_copy_based_d_and_e_labels():
    preview, rows = parse_material_workbook(
        "旧分级.xlsx",
        workbook_bytes([
            ["D消耗不足-有复制", "关键词D"],
            ["E消耗很小-无复制", "关键词E"],
        ]),
    )

    assert preview.valid_count == 2
    assert [row["campaign_name"] for row in rows] == ["D消耗不足", "E消耗很小"]


def test_preview_chunks_every_5000_keywords():
    rows = [["计划A", f"关键词{i}"] for i in range(MAX_KEYWORDS_PER_UNIT + 1)]
    accounts = [{
        "account_id": 1001,
        "login_name": "账户一",
        "landing_url_template": "https://example.com/{userid}",
    }]
    preview = preview_workbook("material.xlsx", workbook_bytes(rows), accounts)
    assert preview.valid_count == MAX_KEYWORDS_PER_UNIT + 1
    assert preview.unit_chunks[0]["unit_count"] == 2
    assert preview.unit_chunks[0]["account_id"] == 1001
    assert preview.target_accounts == accounts
