from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from search_console.report_export import REPORT_HEADERS, build_daily_report_xlsx


def test_daily_report_xlsx_contains_account_settings_and_metrics():
    content = build_daily_report_xlsx([{
        "date": date(2026, 7, 15),
        "account_name": "测试账户",
        "account_id": 81688281,
        "balance": Decimal("188.50"),
        "cost_status": "成本合格",
        "operator_name": "王康",
        "account_type": "二跳账户",
        "page_type": "科普账户",
        "promotion_page": "科普基木鱼",
        "promotion_link": "https://example.com/page",
        "rebate_rate": Decimal("20"),
        "recharge_account": "钱柜A",
        "impressions": 1000,
        "clicks": 100,
        "spend": Decimal("120.00"),
        "uv": 80,
        "copies": 10,
        "adds": 4,
        "cpc": Decimal("1.20"),
        "uv_cost": Decimal("1.50"),
        "copy_cost": Decimal("12.00"),
        "add_cost": Decimal("30.00"),
        "cash_spend": Decimal("100.00"),
        "cash_copy_cost": Decimal("10.00"),
        "cash_add_cost": Decimal("25.00"),
    }])
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook["账户报表"]
    assert tuple(cell.value for cell in sheet[1]) == REPORT_HEADERS
    assert sheet["B2"].value == "测试账户"
    assert sheet["D2"].value == 188.5
    assert sheet["E2"].value == "成本合格"
    assert sheet["K2"].value == 0.2
    assert sheet["J2"].hyperlink.target == "https://example.com/page"
    assert sheet["X2"].value == 10
    assert sheet["Y2"].value == 25


def test_daily_report_xlsx_neutralizes_formula_like_account_text():
    content = build_daily_report_xlsx([{"account_name": "=HYPERLINK(\"bad\")"}])
    sheet = load_workbook(BytesIO(content), data_only=False)["账户报表"]
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value.startswith("'")
