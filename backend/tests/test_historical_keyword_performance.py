from io import BytesIO
from decimal import Decimal

import pytest
from openpyxl import Workbook

from search_console.historical_keyword_performance import (
    HISTORICAL_HEADERS,
    parse_historical_keyword_performance,
)


def workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(HISTORICAL_HEADERS)
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parser_sums_sheets_and_floors_fractional_count_values():
    content = workbook_bytes(
        {
            "科普关键词分析": [["减肥方法", 10, 3, 12.345, 0.2, 1, 0]],
            "软文关键词分析": [["减肥方法", 5, 2, 7.335, 1.9, 2, 1]],
        }
    )

    rows, report = parse_historical_keyword_performance(content)

    assert len(rows) == 1
    assert rows["减肥方法"] == {
        "keyword_text": "减肥方法",
        "impressions": 15,
        "clicks": 5,
        "spend": Decimal("19.68"),
        "uv": 1,
        "copies": 3,
        "adds": 1,
        "source_rows": 2,
    }
    assert report["fractional_values_floored"] == {
        "impressions": 0,
        "clicks": 0,
        "uv": 2,
        "copies": 0,
        "adds": 0,
    }


def test_parser_rejects_wrong_header_order():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["关键词", "点击", "展现", "消费", "UV", "复制", "加粉"])
    sheet.append(["减肥方法", 1, 2, 3, 4, 5, 6])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="表头必须依次"):
        parse_historical_keyword_performance(output.getvalue())


def test_parser_rejects_negative_metrics():
    content = workbook_bytes(
        {"科普关键词分析": [["减肥方法", 10, 3, -0.01, 2, 1, 0]]}
    )

    with pytest.raises(ValueError, match="不能为负数"):
        parse_historical_keyword_performance(content)
