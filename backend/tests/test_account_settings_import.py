import io
import uuid
from decimal import Decimal
import unittest

from openpyxl import load_workbook

from search_console.account_settings_import import (
    HEADERS,
    build_account_settings_template,
    parse_account_settings_workbook,
)
from search_console.models import Account, AccountType


def make_account(
    account_id: int,
    login_name: str,
    *,
    subject: str | None = None,
) -> Account:
    return Account(
        id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        baidu_account_id=account_id,
        login_name=login_name,
        account_subject=subject,
        manager_login_name="manager",
        account_type=AccountType.SECOND_HOP,
        lifecycle_stage="测试期",
        active_keyword_count=1,
        balance=Decimal("0"),
    )


def save_workbook(workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class AccountSettingsImportTest(unittest.TestCase):
    def test_downloaded_template_has_only_requested_columns_and_one_example(self):
        accounts = [make_account(1001, "账户甲"), make_account(1002, "账户乙")]
        content = build_account_settings_template(accounts)
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook["批量设置"]

        self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
        self.assertEqual(sheet.max_column, 8)
        self.assertEqual(sheet.max_row, 2)
        self.assertTrue(sheet["A2"].value.startswith("示例："))
        self.assertNotIn("账户甲", {cell.value for cell in sheet["A"]})
        self.assertNotIn("账户乙", {cell.value for cell in sheet["A"]})

        plan = parse_account_settings_workbook(content, "模板.xlsx", accounts)
        self.assertEqual(plan.invalid_count, 0)
        self.assertEqual(plan.row_count, 0)
        self.assertEqual(plan.changed_count, 0)

    def test_template_import_returns_only_changed_fields(self):
        account = make_account(1001, "账户甲", subject="主体甲")
        content = build_account_settings_template([account])
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["批量设置"]
        sheet["A2"] = "账户甲"
        sheet["B2"] = "王康"
        sheet["G2"] = 12.5

        plan = parse_account_settings_workbook(save_workbook(workbook), "修改.xlsx", [account])
        self.assertEqual(plan.invalid_count, 0)
        self.assertEqual(plan.changed_count, 1)
        self.assertEqual(plan.rows[0].changes, {
            "operator_name": "王康",
            "page_type": "科普账户",
            "promotion_page": "科普基木鱼",
            "promotion_link": "https://example.com/page",
            "rebate_rate": Decimal("12.5"),
            "recharge_account": "示例钱柜账户",
        })

    def test_template_can_match_one_account_by_exact_subject(self):
        account = make_account(1001, "账户甲", subject="主体甲")
        content = build_account_settings_template([account])
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["批量设置"]
        sheet["A2"] = "主体甲"
        sheet["B2"] = "王康"

        plan = parse_account_settings_workbook(save_workbook(workbook), "主体.xlsx", [account])
        self.assertEqual(plan.invalid_count, 0)
        self.assertEqual(plan.matched_count, 1)

    def test_template_import_rejects_invalid_row_atomically(self):
        account = make_account(1001, "账户甲")
        content = build_account_settings_template([account])
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["批量设置"]
        sheet["A2"] = "账户甲"
        sheet["B2"] = "其他人"
        sheet["F2"] = "example.com/page"

        plan = parse_account_settings_workbook(save_workbook(workbook), "错误.xlsx", [account])
        self.assertEqual(plan.invalid_count, 1)
        self.assertEqual(plan.changed_count, 0)
        self.assertFalse(plan.public()["can_import"])
        self.assertIn("operator_name", plan.errors[0]["message"])

    def test_template_import_rejects_duplicate_account_rows(self):
        account = make_account(1001, "账户甲")
        content = build_account_settings_template([account])
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["批量设置"]
        sheet["A2"] = "账户甲"
        sheet.append([cell.value for cell in sheet[2]])

        plan = parse_account_settings_workbook(save_workbook(workbook), "重复.xlsx", [account])
        self.assertEqual(plan.invalid_count, 1)
        self.assertIn("重复出现", plan.errors[0]["message"])
