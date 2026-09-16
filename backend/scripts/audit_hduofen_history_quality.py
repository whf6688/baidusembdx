"""Audit source files, account attribution, keyword attribution, and fact rollups."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import create_engine, text

from search_console.hduofen_history import discover_history_files


SOURCE_LABELS = {
    "visitors": "实时访问详情（UV）",
    "copy_visitors": "复制实时详情（复制）",
    "conversions": "获客助手实时转化详情（加粉）",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--task-id", required=True)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--storage-root", type=Path, required=True)
    parser.add_argument("--imported-profile", type=Path, required=True)
    parser.add_argument("--current-profile", type=Path, required=True)
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    return parser.parse_args()


def style(sheet, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in list(column)[:300]]
        width = min(max(max(map(len, values), default=8) + 2, 10), 55)
        sheet.column_dimensions[column[0].column_letter].width = width


def append_table(sheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style(sheet)


def fetch(connection, sql: str, params: dict[str, object]) -> list[dict[str, object]]:
    return [dict(row) for row in connection.execute(text(sql), params).mappings()]


def profile_source_counts(profile: dict, source: str) -> dict:
    return dict((profile.get("source_counts") or {}).get(source) or {})


def source_file_audit(source_dir: Path, storage_root: Path) -> list[dict[str, object]]:
    stored: dict[tuple[str, str], dict] = {}
    for path in storage_root.rglob("*.json"):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        source_file = Path(str(payload.get("source_file") or ""))
        source_type = path.parts[-4] if len(path.parts) >= 4 else ""
        stored[(source_type, source_file.name)] = {
            "sha256": payload.get("source_sha256"),
            "record_count": payload.get("record_count"),
            "stored_path": str(path),
        }

    rows: list[dict[str, object]] = []
    for item in discover_history_files(source_dir):
        current_hash = hashlib.sha256(item.path.read_bytes()).hexdigest()
        key = (item.source_type, item.path.name)
        previous = stored.get(key)
        status = "未导入" if previous is None else (
            "内容已变化" if previous.get("sha256") != current_hash else "一致"
        )
        rows.append(
            {
                "source_type": item.source_type,
                "source_label": SOURCE_LABELS[item.source_type],
                "file_name": item.path.name,
                "status": status,
                "current_sha256": current_hash,
                "imported_sha256": None if previous is None else previous.get("sha256"),
                "imported_records": None if previous is None else previous.get("record_count"),
            }
        )
    return rows


def main() -> int:
    args = parse_args()
    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL 未配置")
    imported_profile = json.loads(args.imported_profile.read_text(encoding="utf-8"))
    current_profile = json.loads(args.current_profile.read_text(encoding="utf-8"))
    file_rows = source_file_audit(args.source_dir, args.storage_root)
    file_status = Counter(str(row["status"]) for row in file_rows)

    params = {"project_code": args.project_code, "task_id": args.task_id}
    engine = create_engine(database_url, pool_pre_ping=True)
    with engine.connect() as connection:
        task = fetch(
            connection,
            """
            SELECT id::text, status::text, current_node, progress, created_at,
                   heartbeat_at, result, last_error
            FROM search_marketing.background_tasks
            WHERE id = CAST(:task_id AS uuid)
            """,
            params,
        )[0]
        source_rows = fetch(
            connection,
            """
            SELECT event.source_type,
                   count(*) AS records,
                   count(*) FILTER (WHERE event.metric_included) AS included,
                   count(*) FILTER (WHERE NOT event.metric_included) AS excluded,
                   count(*) FILTER (
                       WHERE event.tracking_keyword IS NOT NULL
                   ) AS tracking_keyword_present,
                   count(*) FILTER (
                       WHERE event.metric_included
                         AND event.tracking_keyword IS NOT NULL
                   ) AS included_tracking_keyword_present,
                   count(*) FILTER (
                       WHERE event.metric_included
                         AND event.tracking_keyword IS NOT NULL
                         AND event.material_keyword_id IS NOT NULL
                   ) AS keyword_material_matched,
                   count(DISTINCT event.account_id) FILTER (
                       WHERE event.metric_included
                   ) AS included_accounts,
                   min(event.event_date) AS first_date,
                   max(event.event_date) AS last_date,
                   count(DISTINCT event.metric_key) FILTER (
                       WHERE event.metric_included
                         AND event.source_type = 'visitors'
                   ) AS metric_value
            FROM search_marketing.hduofen_events AS event
            WHERE event.capture_task_id = CAST(:task_id AS uuid)
            GROUP BY event.source_type
            ORDER BY event.source_type
            """,
            params,
        )
        exclusions = fetch(
            connection,
            """
            SELECT source_type, coalesce(exclusion_reason, 'none') AS reason,
                   count(*) AS records,
                   count(DISTINCT raw_data_path) AS files
            FROM search_marketing.hduofen_events
            WHERE capture_task_id = CAST(:task_id AS uuid)
              AND NOT metric_included
            GROUP BY source_type, coalesce(exclusion_reason, 'none')
            ORDER BY records DESC, source_type, reason
            """,
            params,
        )
        account_types = fetch(
            connection,
            """
            SELECT event.source_type, account.account_type::text AS account_type,
                   count(*) AS included_records,
                   count(DISTINCT event.account_id) AS accounts,
                   count(*) FILTER (
                       WHERE event.tracking_keyword IS NOT NULL
                   ) AS tracking_keyword_present,
                   count(*) FILTER (
                       WHERE event.tracking_keyword IS NULL
                   ) AS tracking_keyword_missing
            FROM search_marketing.hduofen_events AS event
            JOIN search_marketing.accounts AS account ON account.id = event.account_id
            WHERE event.capture_task_id = CAST(:task_id AS uuid)
              AND event.metric_included
            GROUP BY event.source_type, account.account_type::text
            ORDER BY event.source_type, account.account_type::text
            """,
            params,
        )
        integrity = fetch(
            connection,
            """
            SELECT
              count(*) AS records,
              count(DISTINCT (event.source_type, event.source_event_id)) AS unique_records,
              count(*) FILTER (WHERE event.metric_included AND event.account_id IS NULL)
                AS included_without_account,
              count(*) FILTER (
                WHERE event.metric_included AND event.tracking_keyword IS NOT NULL
                  AND event.material_keyword_id IS NULL
              ) AS included_keyword_without_material,
              count(*) FILTER (
                WHERE event.metric_included AND event.source_type = 'visitors'
                  AND account.account_type::text = 'SECOND_HOP'
                  AND nullif(btrim(event.tracking_keyword), '') IS NULL
              ) AS invalid_second_hop_visitors,
              count(*) FILTER (
                WHERE event.account_id IS NOT NULL
                  AND event.baidu_account_id IS DISTINCT FROM account.baidu_account_id
              ) AS account_id_mismatches,
              count(*) FILTER (
                WHERE event.material_keyword_id IS NOT NULL
                  AND event.tracking_keyword IS DISTINCT FROM material.keyword_text
              ) AS keyword_text_mismatches,
              count(*) FILTER (
                WHERE event.event_date <> (event.event_at AT TIME ZONE 'Asia/Shanghai')::date
              ) AS timezone_date_mismatches,
              count(DISTINCT event.tracking_keyword) FILTER (
                WHERE event.tracking_keyword IS NOT NULL
              ) AS distinct_tracking_keywords,
              count(DISTINCT event.tracking_keyword) FILTER (
                WHERE event.metric_included AND event.tracking_keyword IS NOT NULL
              ) AS distinct_included_tracking_keywords,
              count(DISTINCT event.material_keyword_id) FILTER (
                WHERE event.metric_included AND event.material_keyword_id IS NOT NULL
              ) AS distinct_material_keywords
            FROM search_marketing.hduofen_events AS event
            LEFT JOIN search_marketing.accounts AS account ON account.id = event.account_id
            LEFT JOIN search_marketing.material_keywords AS material
              ON material.id = event.material_keyword_id
            WHERE event.capture_task_id = CAST(:task_id AS uuid)
            """,
            params,
        )[0]
        fact_checks = fetch(
            connection,
            """
            WITH task_dates AS (
              SELECT DISTINCT event_date
              FROM search_marketing.hduofen_events
              WHERE capture_task_id = CAST(:task_id AS uuid)
            ), account_events AS (
              SELECT event.event_date, event.account_id,
                     count(DISTINCT CASE WHEN event.source_type = 'visitors'
                                         THEN event.metric_key END) AS uv,
                     count(*) FILTER (WHERE event.source_type = 'copy_visitors') AS copies,
                     count(*) FILTER (WHERE event.source_type = 'conversions') AS adds
              FROM search_marketing.hduofen_events AS event
              WHERE event.project_id = (
                    SELECT id FROM search_marketing.projects WHERE code = :project_code
                  )
                AND event.metric_included AND event.account_id IS NOT NULL
                AND event.event_date IN (SELECT event_date FROM task_dates)
              GROUP BY event.event_date, event.account_id
            ), account_compare AS (
              SELECT events.*,
                     coalesce(fact.uv, 0) AS fact_uv,
                     coalesce(fact.copies, 0) AS fact_copies,
                     coalesce(fact.adds, 0) AS fact_adds
              FROM account_events AS events
              LEFT JOIN search_marketing.performance_daily AS fact
                ON fact.report_date = events.event_date
               AND fact.account_id = events.account_id
            ), keyword_events AS (
              SELECT event.event_date, event.account_id,
                     material.campaign_name, material.keyword_text,
                     count(DISTINCT CASE WHEN event.source_type = 'visitors'
                                         THEN event.metric_key END) AS uv,
                     count(*) FILTER (WHERE event.source_type = 'copy_visitors') AS copies,
                     count(*) FILTER (WHERE event.source_type = 'conversions') AS adds
              FROM search_marketing.hduofen_events AS event
              JOIN search_marketing.material_keywords AS material
                ON material.id = event.material_keyword_id
              WHERE event.project_id = (
                    SELECT id FROM search_marketing.projects WHERE code = :project_code
                  )
                AND event.metric_included AND event.account_id IS NOT NULL
                AND event.event_date IN (SELECT event_date FROM task_dates)
              GROUP BY event.event_date, event.account_id,
                       material.campaign_name, material.keyword_text
            ), keyword_compare AS (
              SELECT events.*,
                     coalesce(fact.uv, 0) AS fact_uv,
                     coalesce(fact.copies, 0) AS fact_copies,
                     coalesce(fact.adds, 0) AS fact_adds
              FROM keyword_events AS events
              LEFT JOIN search_marketing.keyword_performance_daily AS fact
                ON fact.report_date = events.event_date
               AND fact.account_id = events.account_id
               AND fact.campaign_name = events.campaign_name
               AND fact.keyword_text = events.keyword_text
            )
            SELECT '账户日汇总' AS check_name, count(*) AS keys,
                   count(*) FILTER (
                     WHERE uv <> fact_uv OR copies <> fact_copies OR adds <> fact_adds
                   ) AS mismatch_keys,
                   coalesce(sum(abs(uv - fact_uv)), 0) AS uv_difference,
                   coalesce(sum(abs(copies - fact_copies)), 0) AS copy_difference,
                   coalesce(sum(abs(adds - fact_adds)), 0) AS add_difference
            FROM account_compare
            UNION ALL
            SELECT '账户关键词日汇总', count(*),
                   count(*) FILTER (
                     WHERE uv <> fact_uv OR copies <> fact_copies OR adds <> fact_adds
                   ),
                   coalesce(sum(abs(uv - fact_uv)), 0),
                   coalesce(sum(abs(copies - fact_copies)), 0),
                   coalesce(sum(abs(adds - fact_adds)), 0)
            FROM keyword_compare
            """,
            params,
        )
        top_excluded_files = fetch(
            connection,
            """
            SELECT source_type, exclusion_reason,
                   regexp_replace(raw_data_path, '^.*/', '') AS file_name,
                   count(*) AS records
            FROM search_marketing.hduofen_events
            WHERE capture_task_id = CAST(:task_id AS uuid)
              AND NOT metric_included
            GROUP BY source_type, exclusion_reason,
                     regexp_replace(raw_data_path, '^.*/', '')
            ORDER BY records DESC, file_name
            LIMIT 200
            """,
            params,
        )
        samples = fetch(
            connection,
            """
            SELECT source_type, event_date, tracking_keyword, exclusion_reason,
                   regexp_replace(raw_data_path, '^.*/', '') AS file_name,
                   source_event_id
            FROM search_marketing.hduofen_events
            WHERE capture_task_id = CAST(:task_id AS uuid)
              AND NOT metric_included
            ORDER BY exclusion_reason, event_date, source_event_id
            LIMIT 500
            """,
            params,
        )

    source_map = {str(row["source_type"]): row for row in source_rows}
    source_comparison: list[dict[str, object]] = []
    account_methods: list[dict[str, object]] = []
    for source in ("visitors", "copy_visitors", "conversions"):
        imported = profile_source_counts(imported_profile, source)
        current = profile_source_counts(current_profile, source)
        database = source_map.get(source, {})
        source_comparison.append(
            {
                "source_type": source,
                "source_label": SOURCE_LABELS[source],
                "imported_files": int(imported.get("files") or 0),
                "current_files": int(current.get("files") or 0),
                "imported_records": int(database.get("records") or 0),
                "current_records": int(current.get("records") or 0),
                "record_delta": int(current.get("records") or 0)
                - int(database.get("records") or 0),
                "imported_included": int(database.get("included") or 0),
                "current_included": int(current.get("included") or 0),
                "included_delta": int(current.get("included") or 0)
                - int(database.get("included") or 0),
            }
        )
        url_total = int(imported.get("account_from_url") or 0)
        url_failed = int(imported.get("custom_id_not_found") or 0) + int(
            imported.get("account_not_found") or 0
        )
        remark_matched = int(imported.get("account_from_remark") or 0)
        remark_failed = int(imported.get("remark_account_not_found") or 0) + int(
            imported.get("missing_account_reference") or 0
        )
        account_methods.append(
            {
                "source_type": source,
                "source_label": SOURCE_LABELS[source],
                "url_total": url_total,
                "url_matched": url_total - url_failed,
                "url_failed": url_failed,
                "url_match_rate_percent": round(
                    (url_total - url_failed) / url_total * 100, 4
                ) if url_total else None,
                "remark_total": remark_matched + remark_failed,
                "remark_matched": remark_matched,
                "remark_failed": remark_failed,
                "remark_match_rate_percent": round(
                    remark_matched / (remark_matched + remark_failed) * 100, 4
                ) if remark_matched + remark_failed else None,
            }
        )

    exclusion_map: Counter[str] = Counter()
    for row in exclusions:
        exclusion_map[str(row["reason"])] += int(row["records"] or 0)
    outside_window = exclusion_map.get("outside_filename_window", 0)
    total_records = int(integrity["records"] or 0)
    in_window_records = total_records - outside_window
    included = sum(int(row.get("included") or 0) for row in source_rows)
    account_matched = included + sum(
        int(row.get("records") or 0)
        for row in exclusions
        if row.get("reason") == "second_hop_missing_tracking_keyword"
    )
    keyword_present = sum(
        int(row.get("included_tracking_keyword_present") or 0)
        for row in source_rows
    )
    keyword_material_matched = sum(
        int(row.get("keyword_material_matched") or 0) for row in source_rows
    )
    summary = {
        "task_id": args.task_id,
        "task_status": task["status"],
        "task_progress": int(task["progress"] or 0),
        "imported_files": int((task.get("result") or {}).get("processed_files") or 0),
        "current_files": len(file_rows),
        "source_files_consistent": int(file_status.get("一致", 0)),
        "source_files_changed": int(file_status.get("内容已变化", 0)),
        "source_files_not_imported": int(file_status.get("未导入", 0)),
        "imported_records": total_records,
        "current_records": int((current_profile.get("counts") or {}).get("records") or 0),
        "current_record_delta": int((current_profile.get("counts") or {}).get("records") or 0)
        - total_records,
        "included_records": included,
        "included_rate_percent": round(included / total_records * 100, 4),
        "included_rate_in_window_percent": round(
            included / in_window_records * 100, 4
        ),
        "account_resolved_records": account_matched,
        "account_resolution_rate_percent": round(
            account_matched / in_window_records * 100, 4
        ),
        "tracking_keyword_present_records": keyword_present,
        "keyword_material_matched_records": keyword_material_matched,
        "keyword_material_match_rate_percent": round(
            keyword_material_matched / keyword_present * 100, 4
        ) if keyword_present else None,
        "distinct_tracking_keywords": int(integrity["distinct_tracking_keywords"] or 0),
        "distinct_included_tracking_keywords": int(
            integrity["distinct_included_tracking_keywords"] or 0
        ),
        "distinct_material_keywords": int(integrity["distinct_material_keywords"] or 0),
        "duplicate_records": total_records - int(integrity["unique_records"] or 0),
        "invalid_second_hop_visitors": int(integrity["invalid_second_hop_visitors"] or 0),
        "included_without_account": int(integrity["included_without_account"] or 0),
        "included_keyword_without_material": int(
            integrity["included_keyword_without_material"] or 0
        ),
        "account_id_mismatches": int(integrity["account_id_mismatches"] or 0),
        "keyword_text_mismatches": int(integrity["keyword_text_mismatches"] or 0),
        "timezone_date_mismatches": int(integrity["timezone_date_mismatches"] or 0),
    }

    issues: list[dict[str, object]] = []
    if summary["source_files_not_imported"] or summary["source_files_changed"]:
        issues.append(
            {
                "severity": "高",
                "issue": "当前源目录与昨天导入快照不一致",
                "evidence": (
                    f"未导入文件 {summary['source_files_not_imported']} 个，"
                    f"内容变化文件 {summary['source_files_changed']} 个，"
                    f"当前比数据库多 {summary['current_record_delta']} 条记录"
                ),
                "impact": "7月13日的UV、复制、加粉可能少计",
                "recommendation": "对新增和变更文件执行增量替换导入，避免整批重导造成重复",
            }
        )
    unresolved = sum(
        exclusion_map.get(reason, 0)
        for reason in (
            "custom_id_not_found",
            "account_not_found",
            "remark_account_not_found",
            "remark_account_ambiguous",
            "missing_account_reference",
        )
    )
    if unresolved:
        issues.append(
            {
                "severity": "高",
                "issue": "账户未匹配记录较多",
                "evidence": f"{unresolved} 条记录无法归属账户，占导入记录 {unresolved / total_records * 100:.2f}%",
                "impact": "对应UV、复制、加粉未进入账户与关键词指标",
                "recommendation": "优先补充自定义ID映射和规范化备注账户名，再对排除记录重算",
            }
        )
    outside = exclusion_map.get("outside_filename_window", 0)
    if outside:
        issues.append(
            {
                "severity": "中",
                "issue": "记录时间超出文件名日期范围",
                "evidence": f"{outside} 条记录被保护性排除",
                "impact": "不会污染正式事实表，但源文件时间口径需确认",
                "recommendation": "核查这些文件的导出时间筛选或事件时间字段",
            }
        )
    second_hop_missing = exclusion_map.get("second_hop_missing_tracking_keyword", 0)
    if second_hop_missing:
        issues.append(
            {
                "severity": "中",
                "issue": "二跳实时访问缺少追踪关键词",
                "evidence": f"{second_hop_missing} 条访问按规则排除，未进入UV指标",
                "impact": "账户和关键词UV都会低于好多粉原始访问量",
                "recommendation": "检查二跳落地页追踪参数，确保访问均携带追踪关键词",
            }
        )
    second_hop_attribution_missing = sum(
        int(row.get("tracking_keyword_missing") or 0)
        for row in account_types
        if row.get("account_type") == "SECOND_HOP"
        and row.get("source_type") in {"copy_visitors", "conversions"}
    )
    if second_hop_attribution_missing:
        issues.append(
            {
                "severity": "中",
                "issue": "部分二跳复制和加粉无法归因到关键词",
                "evidence": f"{second_hop_attribution_missing} 条已归属账户但没有追踪关键词",
                "impact": "账户级复制/加粉正确，关键词级复制/加粉会少计",
                "recommendation": "保留账户指标，同时将无关键词记录列入待归因队列",
            }
        )
    missing_day_count = sum(
        len(item.get("missing_dates") or [])
        for item in (current_profile.get("file_coverage") or {}).values()
    )
    if missing_day_count:
        issues.append(
            {
                "severity": "高",
                "issue": "历史文件日期不连续",
                "evidence": "三类报表均缺少2026-05-18；实时访问另缺2026-07-12",
                "impact": "缺失日期的UV、复制、加粉无法形成完整漏斗",
                "recommendation": "补下载缺失日期文件后再做增量导入",
            }
        )
    if not any(int(row.get("mismatch_keys") or 0) for row in fact_checks):
        issues.append(
            {
                "severity": "通过",
                "issue": "账户与关键词事实表汇总一致",
                "evidence": "账户日及账户关键词日的UV、复制、加粉差异均为0",
                "impact": "数据库汇总链路可信",
                "recommendation": "保持现有唯一键与汇总校验",
            }
        )

    workbook = Workbook()
    overview = workbook.active
    overview.title = "结论摘要"
    append_table(overview, ["指标", "值"], [[key, value] for key, value in summary.items()])
    issue_sheet = workbook.create_sheet("问题清单")
    append_table(
        issue_sheet,
        ["严重程度", "问题", "证据", "影响", "建议"],
        [[row[key] for key in ("severity", "issue", "evidence", "impact", "recommendation")] for row in issues],
    )
    source_sheet = workbook.create_sheet("来源类型核对")
    append_table(
        source_sheet,
        ["来源", "类型", "已导入文件", "当前文件", "已导入记录", "当前记录", "记录差额", "已纳入", "当前可纳入", "可纳入差额"],
        [[row[key] for key in ("source_label", "source_type", "imported_files", "current_files", "imported_records", "current_records", "record_delta", "imported_included", "current_included", "included_delta")] for row in source_comparison],
    )
    exclusion_sheet = workbook.create_sheet("排除原因")
    append_table(
        exclusion_sheet,
        ["来源", "排除原因", "记录数", "涉及文件数"],
        [[SOURCE_LABELS.get(str(row["source_type"]), row["source_type"]), row["reason"], row["records"], row["files"]] for row in exclusions],
    )
    account_sheet = workbook.create_sheet("账户类型与关键词")
    append_table(
        account_sheet,
        ["来源", "账户类型", "纳入记录", "账户数", "有追踪关键词", "无追踪关键词"],
        [[SOURCE_LABELS.get(str(row["source_type"]), row["source_type"]), row["account_type"], row["included_records"], row["accounts"], row["tracking_keyword_present"], row["tracking_keyword_missing"]] for row in account_types],
    )
    method_sheet = workbook.create_sheet("账户匹配方式")
    append_table(
        method_sheet,
        ["来源", "URL有ID", "URL匹配", "URL未匹配", "URL匹配率%", "备注有值", "备注匹配", "备注未匹配", "备注匹配率%"],
        [[row[key] for key in ("source_label", "url_total", "url_matched", "url_failed", "url_match_rate_percent", "remark_total", "remark_matched", "remark_failed", "remark_match_rate_percent")] for row in account_methods],
    )
    fact_sheet = workbook.create_sheet("事实表一致性")
    append_table(
        fact_sheet,
        ["检查", "键数量", "不一致键", "UV差异", "复制差异", "加粉差异"],
        [[row["check_name"], row["keys"], row["mismatch_keys"], row["uv_difference"], row["copy_difference"], row["add_difference"]] for row in fact_checks],
    )
    file_sheet = workbook.create_sheet("源文件差异")
    append_table(
        file_sheet,
        ["来源", "文件名", "状态", "导入记录数", "当前SHA256", "导入SHA256"],
        [[row["source_label"], row["file_name"], row["status"], row["imported_records"], row["current_sha256"], row["imported_sha256"]] for row in file_rows if row["status"] != "一致"],
    )
    top_sheet = workbook.create_sheet("排除最多文件")
    append_table(
        top_sheet,
        ["来源", "排除原因", "文件名", "记录数"],
        [[SOURCE_LABELS.get(str(row["source_type"]), row["source_type"]), row["exclusion_reason"], row["file_name"], row["records"]] for row in top_excluded_files],
    )
    sample_sheet = workbook.create_sheet("排除样本")
    append_table(
        sample_sheet,
        ["来源", "日期", "追踪关键词", "排除原因", "文件名", "事件ID"],
        [[SOURCE_LABELS.get(str(row["source_type"]), row["source_type"]), row["event_date"], row["tracking_keyword"], row["exclusion_reason"], row["file_name"], row["source_event_id"]] for row in samples],
    )
    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_xlsx)

    result = {
        "summary": summary,
        "issues": issues,
        "source_comparison": source_comparison,
        "exclusions": exclusions,
        "account_types": account_types,
        "account_methods": account_methods,
        "fact_checks": fact_checks,
        "file_differences": [row for row in file_rows if row["status"] != "一致"],
        "file_coverage": current_profile.get("file_coverage"),
        "top_excluded_files": top_excluded_files,
    }
    args.output_json.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
