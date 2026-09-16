from __future__ import annotations

import argparse
import json
import os
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy import create_engine, text


EXCEL_DETAIL_LIMIT = 1_000_000


def load_database_url() -> str:
    if os.getenv("DATABASE_URL"):
        return str(os.environ["DATABASE_URL"])
    env_path = Path(__file__).resolve().parents[2] / ".env"
    for line in env_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("DATABASE_URL="):
            value = line.split("=", 1)[1].strip()
            return value.replace("@postgres:", "@127.0.0.1:")
    raise RuntimeError("DATABASE_URL 未配置")


def excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        value = value.replace(tzinfo=None)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def append_rows(sheet, rows: Iterable[Iterable[Any]]) -> int:
    count = 0
    for row in rows:
        sheet.append([excel_value(value) for value in row])
        count += 1
    return count


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export unmatched second-hop keyword report facts."
    )
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--date-from", type=date.fromisoformat, required=True)
    parser.add_argument("--date-to", type=date.fromisoformat, required=True)
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    args = parser.parse_args()
    if args.date_to < args.date_from:
        raise ValueError("结束日期不能早于开始日期")

    params = {
        "project_code": args.project_code,
        "date_from": args.date_from,
        "date_to": args.date_to,
    }
    engine = create_engine(load_database_url(), pool_pre_ping=True)
    workbook = Workbook(write_only=True)

    with engine.connect() as connection:
        summary = dict(
            connection.execute(
                text(
                    """
                    SELECT count(*) AS detail_rows,
                           count(DISTINCT fact.keyword_text) AS keywords,
                           count(DISTINCT fact.account_id) AS accounts,
                           count(DISTINCT fact.campaign_name) AS campaigns,
                           coalesce(sum(fact.impressions), 0) AS impressions,
                           coalesce(sum(fact.clicks), 0) AS clicks,
                           coalesce(sum(fact.spend), 0) AS spend,
                           min(fact.report_date) AS first_date,
                           max(fact.report_date) AS last_date,
                           max(fact.source_watermark) AS source_watermark
                    FROM search_marketing.unmatched_keyword_performance_daily fact
                    JOIN search_marketing.accounts account
                      ON account.id = fact.account_id
                    JOIN search_marketing.projects project
                      ON project.id = account.project_id
                    WHERE project.code = :project_code
                      AND account.account_type::text = 'SECOND_HOP'
                      AND fact.report_date BETWEEN :date_from AND :date_to
                    """
                ),
                params,
            ).one()._mapping
        )
        summary.update(
            {
                "project_code": args.project_code,
                "account_type": "二跳账户",
                "date_from": args.date_from,
                "date_to": args.date_to,
                "match_rule": "去除[已删除]后与物料中心关键词精确匹配",
                "reason": "not_in_material_center",
                "notice": "本文件仅供核对，未匹配关键词未自动加入物料中心。",
            }
        )

        cover = workbook.create_sheet("说明")
        cover.append(["项目", "值"])
        for key, value in summary.items():
            cover.append([key, excel_value(value)])

        keyword_sheet = workbook.create_sheet("关键词汇总")
        keyword_sheet.freeze_panes = "A2"
        keyword_sheet.append(
            [
                "未匹配关键词",
                "账户数",
                "计划数",
                "首次日期",
                "末次日期",
                "展现",
                "点击",
                "消耗",
                "每日记录数",
            ]
        )
        keyword_result = connection.execution_options(stream_results=True).execute(
            text(
                """
                SELECT fact.keyword_text,
                       count(DISTINCT fact.account_id) AS account_count,
                       count(DISTINCT fact.campaign_name) AS campaign_count,
                       min(fact.report_date) AS first_date,
                       max(fact.report_date) AS last_date,
                       sum(fact.impressions) AS impressions,
                       sum(fact.clicks) AS clicks,
                       sum(fact.spend) AS spend,
                       count(*) AS detail_rows
                FROM search_marketing.unmatched_keyword_performance_daily fact
                JOIN search_marketing.accounts account
                  ON account.id = fact.account_id
                JOIN search_marketing.projects project
                  ON project.id = account.project_id
                WHERE project.code = :project_code
                  AND account.account_type::text = 'SECOND_HOP'
                  AND fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.keyword_text
                ORDER BY sum(fact.spend) DESC, fact.keyword_text
                """
            ),
            params,
        ).yield_per(5_000)
        append_rows(keyword_sheet, keyword_result)

        account_keyword_sheet = workbook.create_sheet("账户关键词汇总")
        account_keyword_sheet.freeze_panes = "A2"
        account_keyword_sheet.append(
            [
                "账号主体",
                "账户",
                "账户ID",
                "未匹配关键词",
                "计划数",
                "首次日期",
                "末次日期",
                "展现",
                "点击",
                "消耗",
                "每日记录数",
            ]
        )
        account_keyword_result = connection.execution_options(
            stream_results=True
        ).execute(
            text(
                """
                SELECT account.account_subject, account.login_name,
                       account.baidu_account_id, fact.keyword_text,
                       count(DISTINCT fact.campaign_name) AS campaign_count,
                       min(fact.report_date) AS first_date,
                       max(fact.report_date) AS last_date,
                       sum(fact.impressions) AS impressions,
                       sum(fact.clicks) AS clicks,
                       sum(fact.spend) AS spend,
                       count(*) AS detail_rows
                FROM search_marketing.unmatched_keyword_performance_daily fact
                JOIN search_marketing.accounts account
                  ON account.id = fact.account_id
                JOIN search_marketing.projects project
                  ON project.id = account.project_id
                WHERE project.code = :project_code
                  AND account.account_type::text = 'SECOND_HOP'
                  AND fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY account.account_subject, account.login_name,
                         account.baidu_account_id, fact.keyword_text
                ORDER BY sum(fact.spend) DESC, account.baidu_account_id,
                         fact.keyword_text
                """
            ),
            params,
        ).yield_per(5_000)
        append_rows(account_keyword_sheet, account_keyword_result)

        detail_result = connection.execution_options(stream_results=True).execute(
            text(
                """
                SELECT account.account_subject, account.login_name,
                       account.baidu_account_id, fact.report_date,
                       fact.campaign_name, fact.keyword_text,
                       fact.raw_keyword_text, fact.impressions,
                       fact.clicks, fact.spend, fact.reason,
                       fact.source_watermark
                FROM search_marketing.unmatched_keyword_performance_daily fact
                JOIN search_marketing.accounts account
                  ON account.id = fact.account_id
                JOIN search_marketing.projects project
                  ON project.id = account.project_id
                WHERE project.code = :project_code
                  AND account.account_type::text = 'SECOND_HOP'
                  AND fact.report_date BETWEEN :date_from AND :date_to
                ORDER BY account.baidu_account_id, fact.report_date,
                         fact.campaign_name, fact.keyword_text
                """
            ),
            params,
        ).yield_per(5_000)
        detail_headers = [
            "账号主体",
            "账户",
            "账户ID",
            "日期",
            "计划",
            "规范关键词",
            "原始关键词",
            "展现",
            "点击",
            "消耗",
            "未匹配原因",
            "抓取时间",
        ]
        detail_sheet = None
        detail_count = 0
        detail_sheet_number = 0
        for row in detail_result:
            if detail_sheet is None or detail_count >= EXCEL_DETAIL_LIMIT:
                detail_sheet_number += 1
                detail_sheet = workbook.create_sheet(
                    f"每日明细{detail_sheet_number}"
                )
                detail_sheet.freeze_panes = "A2"
                detail_sheet.append(detail_headers)
                detail_count = 0
            detail_sheet.append([excel_value(value) for value in row])
            detail_count += 1

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_xlsx)
    args.output_json.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=excel_value),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, default=excel_value))


if __name__ == "__main__":
    main()
