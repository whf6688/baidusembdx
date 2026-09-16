"""Import explicit project-scoped Hduofen remark-to-account mappings."""

from __future__ import annotations

import argparse
import hashlib
import json
import uuid
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert

from search_console.db import SessionLocal
from search_console.models import (
    Account,
    AuditEvent,
    HduofenAccountRemarkMapping,
    Project,
)
from search_console.tracking_ingestion import normalize_hduofen_remark


def clean(value) -> str:
    return str(value or "").strip()


def read_rows(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows, ())]
    if "备注/url备注" not in headers or "账户" not in headers:
        workbook.close()
        raise ValueError("工作簿缺少“备注/url备注”或“账户”列")
    remark_index = headers.index("备注/url备注")
    account_index = headers.index("账户")
    result = []
    for values in rows:
        remark = clean(values[remark_index] if remark_index < len(values) else None)
        account = clean(values[account_index] if account_index < len(values) else None)
        if remark or account:
            result.append((remark, account))
    workbook.close()
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--report-path", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    path = Path(args.file)
    rows = read_rows(path)
    source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    normalized = [(normalize_hduofen_remark(remark), remark, account) for remark, account in rows]
    remark_counts = Counter(key for key, _, _ in normalized)
    account_counts = Counter(account for _, _, account in normalized)
    duplicate_remarks = sorted(key for key, count in remark_counts.items() if key and count > 1)
    duplicate_accounts = sorted(key for key, count in account_counts.items() if key and count > 1)
    blank_rows = sum(not key or not account for key, _, account in normalized)

    with SessionLocal() as db:
        project = db.scalar(select(Project).where(Project.code == args.project_code))
        if project is None:
            raise SystemExit(f"项目不存在：{args.project_code}")
        accounts = db.scalars(select(Account).where(Account.project_id == project.id)).all()
        by_name: dict[str, list[Account]] = {}
        for account in accounts:
            by_name.setdefault(account.login_name.strip(), []).append(account)

        matched = []
        unmatched = []
        ambiguous = []
        for key, remark, account_name in normalized:
            candidates = by_name.get(account_name) or []
            if not key or not account_name or not candidates:
                unmatched.append({"remark": remark, "account": account_name})
            elif len(candidates) > 1:
                ambiguous.append({"remark": remark, "account": account_name})
            else:
                matched.append((key, remark, candidates[0]))

        existing = {
            row.normalized_remark: row
            for row in db.scalars(
                select(HduofenAccountRemarkMapping).where(
                    HduofenAccountRemarkMapping.project_id == project.id
                )
            ).all()
        }
        reassigned = sum(
            key in existing and existing[key].account_id != account.id
            for key, _, account in matched
        )
        profile = {
            "generated_at": datetime.now(UTC).isoformat(),
            "source_file": str(path),
            "source_sha256": source_sha256,
            "rows": len(rows),
            "matched": len(matched),
            "unmatched": unmatched,
            "ambiguous": ambiguous,
            "blank_rows": blank_rows,
            "duplicate_remarks": duplicate_remarks,
            "duplicate_accounts": duplicate_accounts,
            "existing_mappings": len(existing),
            "reassigned_mappings": reassigned,
            "apply_requested": args.apply,
        }
        report_path = Path(args.report_path)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(profile, ensure_ascii=False))

        if not args.apply:
            return 0
        if blank_rows or duplicate_remarks or duplicate_accounts or unmatched or ambiguous:
            raise SystemExit("好多粉备注映射预检未通过，已拒绝写库")

        now = datetime.now(UTC)
        values = [
            {
                "id": uuid.uuid4(),
                "project_id": project.id,
                "remark": remark,
                "normalized_remark": key,
                "account_id": account.id,
                "source_file": str(path),
                "source_sha256": source_sha256,
                "created_at": now,
                "updated_at": now,
            }
            for key, remark, account in matched
        ]
        statement = pg_insert(HduofenAccountRemarkMapping).values(values)
        excluded = statement.excluded
        db.execute(
            statement.on_conflict_do_update(
                index_elements=[
                    HduofenAccountRemarkMapping.project_id,
                    HduofenAccountRemarkMapping.normalized_remark,
                ],
                set_={
                    "remark": excluded.remark,
                    "account_id": excluded.account_id,
                    "source_file": excluded.source_file,
                    "source_sha256": excluded.source_sha256,
                    "updated_at": excluded.updated_at,
                },
            )
        )
        db.add(
            AuditEvent(
                project_id=project.id,
                actor="hduofen-remark-mapping-importer",
                action="hduofen.remark_mapping.import",
                target_type="project",
                target_id=str(project.id),
                summary="导入好多粉备注账户映射",
                details={
                    "source_file": str(path),
                    "source_sha256": source_sha256,
                    "mappings": len(values),
                    "reassigned": reassigned,
                },
            )
        )
        db.commit()
        print(json.dumps({"imported": len(values), "reassigned": reassigned}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
