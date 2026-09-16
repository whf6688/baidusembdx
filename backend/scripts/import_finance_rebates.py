"""Import account rebates from the first three account sheets of a finance workbook.

The finance value is a multiplier in the form ``1 + rebate percentage``. Exact
duplicate account names use the bottom-most occurrence; sheets are processed in
workbook order, so a later sheet also takes precedence over an earlier sheet.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from search_console.db import SessionLocal
from search_console.models import Account, AuditEvent, Project


def clean(value: object) -> str:
    return str(value or "").strip()


def read_latest_rebates(path: Path) -> tuple[dict[str, dict], list[dict], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    latest: dict[str, dict] = {}
    sheet_stats: list[dict] = []
    invalid_source_rows = 0
    for sheet in workbook.worksheets[:3]:
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows, ())]
        if "账户" not in headers or "返点" not in headers:
            workbook.close()
            raise ValueError(f"工作表“{sheet.title}”缺少账户或返点列")
        account_index = headers.index("账户")
        rebate_index = headers.index("返点")
        source_rows = 0
        names: set[str] = set()
        for row_number, values in enumerate(rows, start=2):
            if account_index >= len(values) or rebate_index >= len(values):
                continue
            account_name = clean(values[account_index])
            raw_value = values[rebate_index]
            if not account_name or isinstance(raw_value, bool) or not isinstance(raw_value, (int, float, Decimal)):
                continue
            multiplier = Decimal(str(raw_value))
            rebate_rate = ((multiplier - Decimal("1")) * Decimal("100")).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP
            )
            source_rows += 1
            names.add(account_name)
            if multiplier < 1 or multiplier > 2:
                invalid_source_rows += 1
            latest[account_name] = {
                "sheet": sheet.title,
                "row": row_number,
                "multiplier": str(multiplier),
                "rebate_rate": rebate_rate,
            }
        sheet_stats.append({
            "sheet": sheet.title,
            "source_rows": source_rows,
            "unique_accounts": len(names),
            "duplicate_rows": source_rows - len(names),
        })
    workbook.close()
    return latest, sheet_stats, invalid_source_rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--project-code", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    path = Path(args.file)
    latest, sheet_stats, invalid_source_rows = read_latest_rebates(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    with SessionLocal() as db:
        project = db.scalar(select(Project).where(Project.code == args.project_code))
        if project is None:
            raise ValueError(f"项目不存在：{args.project_code}")
        accounts = db.scalars(select(Account).where(
            Account.project_id == project.id,
            Account.is_active.is_(True),
        )).all()
        by_login = {account.login_name: account for account in accounts}
        matched = [(by_login[name], data) for name, data in latest.items() if name in by_login]
        unmatched = [name for name in latest if name not in by_login]
        invalid_matched = [
            account.login_name
            for account, data in matched
            if Decimal(data["multiplier"]) < 1 or Decimal(data["multiplier"]) > 2
        ]
        if invalid_matched:
            raise ValueError(f"当前项目存在非法返点倍数：{invalid_matched[:20]}")
        changed = [
            (account, data) for account, data in matched
            if account.rebate_rate != data["rebate_rate"]
        ]

        summary = {
            "mode": "apply" if args.apply else "dry-run",
            "project_code": project.code,
            "sheets": sheet_stats,
            "source_unique_accounts": len(latest),
            "exact_matches": len(matched),
            "changed_accounts": len(changed),
            "unmatched_accounts": len(unmatched),
            "unmatched_login_names": unmatched[:50],
            "invalid_source_rows": invalid_source_rows,
            "invalid_matched_accounts": len(invalid_matched),
            "duplicate_rule": "bottom_most_wins",
            "conversion_rule": "(multiplier - 1) * 100",
            "sha256": digest,
        }

        if args.apply:
            for account, data in changed:
                account.rebate_rate = data["rebate_rate"]
            db.add(AuditEvent(
                project_id=project.id,
                actor="codex-finance-import",
                action="account.finance_rebate.update",
                target_type="account",
                summary="按财务工作簿最下方记录批量更新账户返点",
                details=summary,
            ))
            db.commit()

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
