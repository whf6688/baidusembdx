"""Update existing project accounts from the approved account workbook.

Rows are matched by exact Baidu login name. When a login name appears more than
once in the workbook, only its first occurrence is used. The workbook lifecycle
column is intentionally ignored because lifecycle is computed by the platform.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from search_console.db import SessionLocal
from search_console.models import Account, AccountCategory, AccountType, AuditEvent, Project


ALLOWED_PAGE_TYPES = {"科普账户", "软文账户"}
ALLOWED_PROMOTION_PAGES = {"科普基木鱼", "科普全文", "精华帖", "中医论坛", "中医秘方", "快瘦汤"}
ALLOWED_OPERATORS = {"王康", "王聪"}
PAGE_TYPE_ALIASES = {"软文页面": "软文账户"}


def clean(value: object) -> str:
    return str(value or "").strip()


def read_first_rows(path: Path) -> tuple[
    dict[str, dict[str, str | None]], int, int, int, list[str], int, bool, bool, bool, bool
]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [clean(value) for value in raw_headers]
    positions: dict[str, list[int]] = defaultdict(list)
    for index, header in enumerate(headers):
        positions[header].append(index)
    required = {"账户", "账户类型", "页面类型", "推广页面", "运营"}
    missing = sorted(required - positions.keys())
    if missing:
        workbook.close()
        raise ValueError(f"工作簿表头不兼容，缺少字段：{missing}")
    account_index = positions["账户"][0]
    category_index = (
        positions["账户分类"][0]
        if positions.get("账户分类")
        else positions["运营"][0] if len(positions["运营"]) >= 2 else None
    )
    operator_index = positions["运营"][-1]
    account_type_index = positions["账户类型"][0]
    page_type_index = positions["页面类型"][0]
    promotion_page_index = positions["推广页面"][0]
    recharge_account_index = positions.get("钱柜账户", [None])[0]
    promotion_link_index = positions.get("推广链接", [None])[0]
    rebate_rate_index = positions.get("返点", [None])[0]
    used_indexes = {
        account_index, operator_index, account_type_index, page_type_index,
        promotion_page_index,
    }
    if category_index is not None:
        used_indexes.add(category_index)
    if recharge_account_index is not None:
        used_indexes.add(recharge_account_index)
    if promotion_link_index is not None:
        used_indexes.add(promotion_link_index)
    if rebate_rate_index is not None:
        used_indexes.add(rebate_rate_index)
    ignored_columns = [
        header or f"未命名列{index + 1}"
        for index, header in enumerate(headers)
        if index not in used_indexes
    ]
    first_rows: dict[str, dict[str, str | None]] = {}
    data_rows = 0
    duplicate_rows = 0
    conflicting_duplicate_rows = 0
    normalized_page_types = 0
    for values in sheet.iter_rows(min_row=2, values_only=True):
        data_rows += 1
        login_name = clean(values[account_index] if len(values) > account_index else None)
        if not login_name:
            continue
        page_type = clean(values[page_type_index] if len(values) > page_type_index else None)
        normalized_page_type = PAGE_TYPE_ALIASES.get(page_type, page_type)
        row_data = {
            "category_name": (
                clean(values[category_index] if len(values) > category_index else None)
                if category_index is not None else None
            ),
            "account_type": clean(values[account_type_index] if len(values) > account_type_index else None),
            "page_type": normalized_page_type,
            "promotion_page": clean(values[promotion_page_index] if len(values) > promotion_page_index else None),
            "operator_name": clean(values[operator_index] if len(values) > operator_index else None),
            "recharge_account": (
                clean(values[recharge_account_index] if len(values) > recharge_account_index else None)
                if recharge_account_index is not None else None
            ),
            "promotion_link": (
                clean(values[promotion_link_index] if len(values) > promotion_link_index else None) or None
                if promotion_link_index is not None else None
            ),
            "rebate_rate": (
                clean(values[rebate_rate_index] if len(values) > rebate_rate_index else None) or None
                if rebate_rate_index is not None else None
            ),
        }
        if login_name in first_rows:
            duplicate_rows += 1
            if row_data != first_rows[login_name]:
                conflicting_duplicate_rows += 1
            continue
        if normalized_page_type != page_type:
            normalized_page_types += 1
        first_rows[login_name] = row_data
    workbook.close()
    return (
        first_rows, data_rows, duplicate_rows, conflicting_duplicate_rows,
        ignored_columns, normalized_page_types, category_index is not None,
        recharge_account_index is not None, promotion_link_index is not None,
        rebate_rate_index is not None,
    )


def validate_rows(rows: dict[str, dict[str, str | None]]) -> None:
    account_types = {item.value for item in AccountType}
    errors: list[str] = []
    for login_name, row in rows.items():
        if row["account_type"] not in account_types:
            errors.append(f"{login_name}: 未知账户类型 {row['account_type']}")
        if row["page_type"] not in ALLOWED_PAGE_TYPES:
            errors.append(f"{login_name}: 未知页面类型 {row['page_type']}")
        if row["promotion_page"] not in ALLOWED_PROMOTION_PAGES:
            errors.append(f"{login_name}: 未知推广页面 {row['promotion_page']}")
        if row["operator_name"] not in ALLOWED_OPERATORS:
            errors.append(f"{login_name}: 未知运营 {row['operator_name']}")
        if row["category_name"] is not None and not row["category_name"]:
            errors.append(f"{login_name}: 账户分类为空")
        if row["promotion_link"] and not str(row["promotion_link"]).startswith(("http://", "https://")):
            errors.append(f"{login_name}: 推广链接必须以 http:// 或 https:// 开头")
        if row["rebate_rate"] is not None:
            try:
                rebate_rate = Decimal(str(row["rebate_rate"]))
            except Exception:
                errors.append(f"{login_name}: 返点不是有效数字")
            else:
                if rebate_rate < 0 or rebate_rate > 100:
                    errors.append(f"{login_name}: 返点必须在 0 至 100 之间")
    if errors:
        raise ValueError("工作簿字段校验失败：" + "；".join(errors[:20]))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--project-code", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    path = Path(args.file)
    (
        first_rows, data_rows, duplicate_rows, conflicting_duplicate_rows,
        ignored_columns, normalized_page_types, category_column_present,
        recharge_account_column_present, promotion_link_column_present,
        rebate_rate_column_present,
    ) = read_first_rows(path)
    validate_rows(first_rows)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    with SessionLocal() as db:
        project = db.scalar(select(Project).where(Project.code == args.project_code))
        if project is None:
            raise ValueError(f"项目不存在：{args.project_code}")

        accounts = db.scalars(select(Account).where(Account.project_id == project.id)).all()
        by_login: dict[str, list[Account]] = defaultdict(list)
        for account in accounts:
            by_login[account.login_name].append(account)

        matched: list[tuple[Account, dict[str, str]]] = []
        unmatched: list[str] = []
        ambiguous: list[str] = []
        for login_name, row in first_rows.items():
            candidates = by_login.get(login_name, [])
            if len(candidates) == 1:
                matched.append((candidates[0], row))
            elif not candidates:
                unmatched.append(login_name)
            elif len(candidates) > 1:
                ambiguous.append(login_name)

        changed = [
            (account, row)
            for account, row in matched
            if account.account_type.value != row["account_type"]
            or account.operator_name != row["operator_name"]
            or (
                row["category_name"] is not None
                and (account.category.name if account.category else "") != row["category_name"]
            )
            or account.page_type != row["page_type"]
            or account.promotion_page != row["promotion_page"]
            or (
                recharge_account_column_present
                and account.recharge_account != row["recharge_account"]
            )
            or (
                row["promotion_link"] is not None
                and account.landing_url_template != row["promotion_link"]
            )
            or (
                row["rebate_rate"] is not None
                and account.rebate_rate != Decimal(str(row["rebate_rate"]))
            )
        ]
        summary = {
            "mode": "apply" if args.apply else "dry-run",
            "project_code": project.code,
            "workbook_rows": data_rows,
            "unique_accounts_first_row": len(first_rows),
            "duplicate_rows_ignored": duplicate_rows,
            "conflicting_duplicate_rows": conflicting_duplicate_rows,
            "exact_matches": len(matched),
            "changed_accounts": len(changed),
            "unmatched_accounts": len(unmatched),
            "unmatched_login_names": unmatched[:50],
            "ambiguous_accounts": len(ambiguous),
            "ambiguous_login_names": ambiguous[:50],
            "lifecycle_column_ignored": True,
            "ignored_columns": ignored_columns,
            "normalized_page_type_rows": normalized_page_types,
            "category_column_present": category_column_present,
            "recharge_account_column_present": recharge_account_column_present,
            "promotion_link_column_present": promotion_link_column_present,
            "rebate_rate_column_present": rebate_rate_column_present,
            "rebate_format": "percentage",
            "sha256": digest,
        }

        if args.apply:
            category_names = {
                row["category_name"] for _, row in changed
                if row["category_name"] is not None
            }
            categories = {
                category.name: category
                for category in db.scalars(
                    select(AccountCategory).where(
                        AccountCategory.project_id == project.id,
                        AccountCategory.name.in_(category_names),
                    )
                ).all()
            }
            for name in sorted(category_names - categories.keys()):
                category = AccountCategory(project_id=project.id, name=name)
                db.add(category)
                db.flush()
                categories[name] = category

            for account, row in changed:
                if row["category_name"] is not None:
                    account.category_id = categories[row["category_name"]].id
                account.account_type = AccountType(row["account_type"])
                account.page_type = row["page_type"]
                account.promotion_page = row["promotion_page"]
                account.operator_name = row["operator_name"]
                if recharge_account_column_present:
                    account.recharge_account = row["recharge_account"]
                if row["promotion_link"] is not None:
                    account.landing_url_template = row["promotion_link"]
                if row["rebate_rate"] is not None:
                    account.rebate_rate = Decimal(str(row["rebate_rate"]))

            db.add(AuditEvent(
                project_id=project.id,
                actor="codex-workbook-import",
                action="account.workbook.update",
                target_type="account",
                summary="按工作簿首条记录批量更新账户业务信息",
                details=summary,
            ))
            db.commit()

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
