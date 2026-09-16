"""Import non-gibberish Hduofen unmatched keywords into the project material library."""

from __future__ import annotations

import argparse
import hashlib
import json
import uuid
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert

from search_console.db import SessionLocal
from search_console.keyword_encoding import KEYWORD_ENCODING_VERSION, encode_keyword_utf8
from search_console.models import (
    AuditEvent,
    BackgroundTask,
    HduofenEvent,
    Material,
    MaterialKeyword,
    MaterialVersion,
    Project,
)
from search_console.tracking_ingestion import (
    _aggregate_keyword_metrics,
    is_gibberish_tracking_keyword,
    normalize_tracking_keyword,
)


MATERIAL_NAME = "项目公共物料库"
CAMPAIGN_NAME = "好多粉未匹配补充"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--task-id", required=True)
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--report-path", required=True)
    parser.add_argument("--apply", action="store_true")
    return parser.parse_args()


def read_keywords(path: Path) -> tuple[list[str], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "未匹配关键词汇总" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("工作簿缺少“未匹配关键词汇总”")
    sheet = workbook["未匹配关键词汇总"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    raw_index = headers.index("追踪关键词")
    normalized_index = headers.index("归一化关键词")
    result: list[str] = []
    gibberish: list[str] = []
    seen: set[str] = set()
    for values in rows:
        raw = str(values[raw_index] or "").strip()
        normalized = normalize_tracking_keyword(str(values[normalized_index] or raw))
        if not normalized:
            continue
        if is_gibberish_tracking_keyword(raw) or is_gibberish_tracking_keyword(normalized):
            gibberish.append(raw)
            continue
        key = normalized.casefold()
        if key in seen:
            continue
        seen.add(key)
        result.append(normalized)
    if "疑似乱码关键词" in workbook.sheetnames:
        gibberish_sheet = workbook["疑似乱码关键词"]
        gibberish_rows = gibberish_sheet.iter_rows(values_only=True)
        gibberish_headers = [str(value or "").strip() for value in next(gibberish_rows, ())]
        if "追踪关键词" in gibberish_headers:
            gibberish_index = gibberish_headers.index("追踪关键词")
            for values in gibberish_rows:
                value = str(values[gibberish_index] or "").strip()
                if value:
                    gibberish.append(value)
    workbook.close()
    return result, list(dict.fromkeys(gibberish))


def main() -> int:
    args = parse_args()
    path = Path(args.file)
    keywords, gibberish = read_keywords(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    task_id = uuid.UUID(args.task_id)

    with SessionLocal() as db:
        project = db.scalar(select(Project).where(Project.code == args.project_code))
        task = db.get(BackgroundTask, task_id)
        if project is None or task is None or task.project_id != project.id:
            raise RuntimeError("项目或好多粉任务不存在")
        existing_rows = db.scalars(
            select(MaterialKeyword).where(MaterialKeyword.project_id == project.id)
        ).all()
        existing_normalized = {
            normalize_tracking_keyword(row.keyword_text).casefold() for row in existing_rows
        }
        new_keywords = [
            keyword for keyword in keywords if keyword.casefold() not in existing_normalized
        ]
        skipped_existing = len(keywords) - len(new_keywords)

        raw_keywords = db.scalars(
            select(HduofenEvent.tracking_keyword)
            .where(
                HduofenEvent.capture_task_id == task_id,
                HduofenEvent.tracking_keyword.is_not(None),
            )
            .distinct()
        ).all()
        target_keys = {keyword.casefold() for keyword in keywords}
        eligible_raw_keywords = [
            raw
            for raw in raw_keywords
            if not is_gibberish_tracking_keyword(raw)
            and normalize_tracking_keyword(raw).casefold() in target_keys
        ]
        preview = {
            "source_file": str(path),
            "source_sha256": digest,
            "source_non_gibberish_keywords": len(keywords),
            "source_gibberish_skipped": len(gibberish),
            "new_material_keywords": len(new_keywords),
            "existing_material_keywords": skipped_existing,
            "eligible_tracking_keyword_variants": len(eligible_raw_keywords),
            "apply": args.apply,
        }
        report_path = Path(args.report_path)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(preview, ensure_ascii=False, indent=2), encoding="utf-8")
        if not args.apply:
            db.rollback()
            print(json.dumps(preview, ensure_ascii=False, indent=2))
            return 0

        material = db.scalar(
            select(Material).where(
                Material.project_id == project.id,
                Material.name == MATERIAL_NAME,
            )
        )
        if material is None:
            material = Material(project_id=project.id, name=MATERIAL_NAME, current_version=1)
            db.add(material)
            db.flush()
        elif new_keywords:
            material.current_version += 1

        version = None
        if new_keywords:
            version = MaterialVersion(
                material_id=material.id,
                version=material.current_version,
                file_path=str(path),
                sha256=digest,
                row_count=len(new_keywords),
                summary={
                    "source_non_gibberish_keywords": len(keywords),
                    "gibberish_skipped": len(gibberish),
                    "existing_skipped": skipped_existing,
                },
                created_by="hduofen-unmatched-material-importer",
            )
            db.add(version)
            db.flush()
            values = [
                {
                    "id": uuid.uuid4(),
                    "project_id": project.id,
                    "material_version_id": version.id,
                    "row_number": row_number,
                    "campaign_name": CAMPAIGN_NAME,
                    "keyword_text": keyword,
                    "keyword_utf8_encoded": encode_keyword_utf8(keyword),
                    "keyword_encoding_version": KEYWORD_ENCODING_VERSION,
                    "is_blacklisted": False,
                }
                for row_number, keyword in enumerate(new_keywords, start=2)
            ]
            db.execute(
                pg_insert(MaterialKeyword)
                .values(values)
                .on_conflict_do_nothing(
                    index_elements=[MaterialKeyword.project_id, MaterialKeyword.keyword_text]
                )
            )
            db.flush()

        all_materials = db.scalars(
            select(MaterialKeyword).where(MaterialKeyword.project_id == project.id)
        ).all()
        by_normalized: dict[str, list[MaterialKeyword]] = defaultdict(list)
        for row in all_materials:
            by_normalized[normalize_tracking_keyword(row.keyword_text).casefold()].append(row)

        linked_events = 0
        ambiguous_variants: list[str] = []
        affected_dates = set()
        for raw in eligible_raw_keywords:
            candidates = by_normalized.get(normalize_tracking_keyword(raw).casefold()) or []
            exact = next((row for row in candidates if row.keyword_text == raw), None)
            material_keyword = exact or (candidates[0] if len(candidates) == 1 else None)
            if material_keyword is None:
                ambiguous_variants.append(raw)
                continue
            dates = db.scalars(
                select(HduofenEvent.event_date).where(
                    HduofenEvent.capture_task_id == task_id,
                    HduofenEvent.tracking_keyword == raw,
                    HduofenEvent.account_id.is_not(None),
                    HduofenEvent.metric_included.is_(True),
                ).distinct()
            ).all()
            affected_dates.update(dates)
            linked_events += db.execute(
                update(HduofenEvent)
                .where(
                    HduofenEvent.capture_task_id == task_id,
                    HduofenEvent.tracking_keyword == raw,
                )
                .values(material_keyword_id=material_keyword.id)
            ).rowcount

        watermark = datetime.now(UTC)
        _aggregate_keyword_metrics(db, project.id, affected_dates, watermark)
        details = {
            **preview,
            "linked_events": linked_events,
            "affected_dates": len(affected_dates),
            "ambiguous_variants": ambiguous_variants,
        }
        db.add(
            AuditEvent(
                project_id=project.id,
                actor="hduofen-unmatched-material-importer",
                action="hduofen.unmatched_keyword.material_import",
                target_type="background_task",
                target_id=str(task_id),
                summary="将好多粉非乱码未匹配关键词加入物料中心并重新关联",
                details=details,
            )
        )
        db.commit()
        report_path.write_text(json.dumps(details, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(details, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
