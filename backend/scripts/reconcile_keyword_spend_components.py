"""Reconcile second-hop account spend into matched, unmatched, and other spend."""

from __future__ import annotations

import argparse
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import create_engine, text


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--date-from", type=date.fromisoformat, required=True)
    parser.add_argument("--date-to", type=date.fromisoformat, required=True)
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    return parser.parse_args()


def database_url() -> str:
    value = os.getenv("DATABASE_URL")
    if not value:
        raise RuntimeError("DATABASE_URL 未配置")
    return value


def decimal(value: object) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def style(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in list(column)[:200]]
        width = min(max(max(map(len, values), default=8) + 2, 10), 42)
        sheet.column_dimensions[column[0].column_letter].width = width


def main() -> int:
    args = parse_args()
    params = {
        "project_code": args.project_code,
        "date_from": args.date_from,
        "date_to": args.date_to,
    }
    statement = text(
        """
        WITH scoped_accounts AS (
            SELECT account.id, account.baidu_account_id, account.login_name,
                   account.account_subject
            FROM search_marketing.accounts AS account
            JOIN search_marketing.projects AS project
              ON project.id = account.project_id
            WHERE project.code = :project_code
              AND account.is_active
              AND account.account_type::text = 'SECOND_HOP'
        ), account_spend AS (
            SELECT fact.account_id, sum(fact.spend) AS spend
            FROM search_marketing.performance_daily AS fact
            JOIN scoped_accounts AS account ON account.id = fact.account_id
            WHERE fact.report_date BETWEEN :date_from AND :date_to
            GROUP BY fact.account_id
        ), matched_spend AS (
            SELECT fact.account_id, sum(fact.spend) AS spend
            FROM search_marketing.keyword_performance_daily AS fact
            JOIN scoped_accounts AS account ON account.id = fact.account_id
            WHERE fact.report_date BETWEEN :date_from AND :date_to
            GROUP BY fact.account_id
        ), unmatched_spend AS (
            SELECT fact.account_id, sum(fact.spend) AS spend
            FROM search_marketing.unmatched_keyword_performance_daily AS fact
            JOIN scoped_accounts AS account ON account.id = fact.account_id
            WHERE fact.report_date BETWEEN :date_from AND :date_to
            GROUP BY fact.account_id
        )
        SELECT account.account_subject, account.login_name,
               account.baidu_account_id,
               coalesce(account_spend.spend, 0) AS account_spend,
               coalesce(matched_spend.spend, 0) AS matched_spend,
               coalesce(unmatched_spend.spend, 0) AS unmatched_spend,
               coalesce(account_spend.spend, 0)
                 - coalesce(matched_spend.spend, 0)
                 - coalesce(unmatched_spend.spend, 0) AS other_spend
        FROM scoped_accounts AS account
        LEFT JOIN account_spend ON account_spend.account_id = account.id
        LEFT JOIN matched_spend ON matched_spend.account_id = account.id
        LEFT JOIN unmatched_spend ON unmatched_spend.account_id = account.id
        ORDER BY abs(
            coalesce(account_spend.spend, 0)
              - coalesce(matched_spend.spend, 0)
              - coalesce(unmatched_spend.spend, 0)
        ) DESC, account.baidu_account_id
        """
    )
    engine = create_engine(database_url(), pool_pre_ping=True)
    with engine.connect() as connection:
        rows = [dict(row) for row in connection.execute(statement, params).mappings()]
    for row in rows:
        for field in ("account_spend", "matched_spend", "unmatched_spend", "other_spend"):
            row[field] = decimal(row[field])

    totals = {
        field: sum((row[field] for row in rows), Decimal("0"))
        for field in ("account_spend", "matched_spend", "unmatched_spend", "other_spend")
    }
    summary = {
        "project_code": args.project_code,
        "account_type": "二跳账户",
        "date_from": args.date_from.isoformat(),
        "date_to": args.date_to.isoformat(),
        "accounts": len(rows),
        **{key: float(value) for key, value in totals.items()},
        "matched_plus_unmatched": float(totals["matched_spend"] + totals["unmatched_spend"]),
        "classified_coverage_percent": float(
            (
                (totals["matched_spend"] + totals["unmatched_spend"])
                / totals["account_spend"]
                * Decimal("100")
            ).quantize(Decimal("0.0001"))
        ),
        "accounts_with_other_spend": sum(
            1 for row in rows if abs(row["other_spend"]) > Decimal("0.01")
        ),
    }

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "核对摘要"
    summary_sheet.append(["项目", "值"])
    for key, value in summary.items():
        summary_sheet.append([key, value])
    style(summary_sheet)

    detail = workbook.create_sheet("按账户分解")
    detail.append(
        ["账号主体", "账户", "账户ID", "账户消耗", "已匹配关键词消费", "未匹配关键词消费", "其他消费"]
    )
    for row in rows:
        detail.append(
            [
                row["account_subject"],
                row["login_name"],
                row["baidu_account_id"],
                float(row["account_spend"]),
                float(row["matched_spend"]),
                float(row["unmatched_spend"]),
                float(row["other_spend"]),
            ]
        )
    style(detail)
    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_xlsx)
    args.output_json.write_text(
        json.dumps(
            {
                **summary,
                "top_other_spend_accounts": [
                    {
                        "account_id": row["baidu_account_id"],
                        "account": row["login_name"],
                        "other_spend": float(row["other_spend"]),
                    }
                    for row in rows[:20]
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
