"""Import explicit account lifecycle overrides from a two-column workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from search_console.account_lifecycle import MANUAL_LIFECYCLE_STAGES
from search_console.api import refresh_project_account_lifecycles
from search_console.db import SessionLocal
from search_console.models import Account, AuditEvent, Project


AUTO = "自动判断"


def clean(value: object) -> str:
    return str(value or "").strip()


def read_lifecycles(path: Path) -> tuple[dict[str, str], int, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows, ())]
    if "账户" not in headers or "生命周期" not in headers:
        workbook.close()
        raise ValueError("工作簿必须包含账户和生命周期列")
    account_index = headers.index("账户")
    lifecycle_index = headers.index("生命周期")
    latest: dict[str, str] = {}
    source_rows = 0
    duplicate_rows = 0
    allowed = set(MANUAL_LIFECYCLE_STAGES) | {AUTO}
    for values in rows:
        account_name = clean(values[account_index] if len(values) > account_index else None)
        lifecycle = clean(values[lifecycle_index] if len(values) > lifecycle_index else None)
        if not account_name and not lifecycle:
            continue
        if not account_name or lifecycle not in allowed:
            workbook.close()
            raise ValueError(f"生命周期行无效：账户={account_name!r}，生命周期={lifecycle!r}")
        source_rows += 1
        if account_name in latest:
            duplicate_rows += 1
        latest[account_name] = lifecycle
    workbook.close()
    return latest, source_rows, duplicate_rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--project-code", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    path = Path(args.file)
    lifecycles, source_rows, duplicate_rows = read_lifecycles(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    with SessionLocal() as db:
        project = db.scalar(select(Project).where(Project.code == args.project_code))
        if project is None:
            raise ValueError(f"项目不存在：{args.project_code}")
        accounts = db.scalars(select(Account).where(
            Account.project_id == project.id,
            Account.is_active.is_(True),
        )).all()
        by_login = {account.login_name: account for account in accounts}
        matched = [(by_login[name], stage) for name, stage in lifecycles.items() if name in by_login]
        unmatched = [name for name in lifecycles if name not in by_login]
        changed = [
            (account, stage) for account, stage in matched
            if account.lifecycle_override != (None if stage == AUTO else stage)
        ]
        summary = {
            "mode": "apply" if args.apply else "dry-run",
            "project_code": project.code,
            "source_rows": source_rows,
            "unique_accounts": len(lifecycles),
            "duplicate_rows": duplicate_rows,
            "exact_matches": len(matched),
            "changed_accounts": len(changed),
            "unmatched_accounts": len(unmatched),
            "unmatched_login_names": unmatched[:50],
            "stage_counts": {
                stage: sum(value == stage for value in lifecycles.values())
                for stage in sorted(set(lifecycles.values()))
            },
            "sha256": digest,
        }
        if args.apply:
            for account, stage in changed:
                account.lifecycle_override = None if stage == AUTO else stage
            refresh_project_account_lifecycles(db, project.id, commit=False)
            db.add(AuditEvent(
                project_id=project.id,
                actor="codex-lifecycle-import",
                action="account.lifecycle_override.update",
                target_type="account",
                summary="按工作簿批量设置账户生命周期人工覆盖",
                details=summary,
            ))
            db.commit()
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
