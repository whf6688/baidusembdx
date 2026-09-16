from __future__ import annotations

import argparse
import json
import os
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import create_engine, text


MONEY_TOLERANCE = Decimal("0.01")


def load_database_url() -> str:
    if os.getenv("DATABASE_URL"):
        return str(os.environ["DATABASE_URL"])
    env_path = Path(__file__).resolve().parents[2] / ".env"
    for line in env_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("DATABASE_URL="):
            value = line.split("=", 1)[1].strip()
            return value.replace("@postgres:", "@127.0.0.1:")
    raise RuntimeError("DATABASE_URL 未配置")


def as_decimal(value: Any) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def fetch_rows(connection, statement: str, params: dict[str, Any]) -> list[dict]:
    return [
        {key: json_value(value) for key, value in row._mapping.items()}
        for row in connection.execute(text(statement), params)
    ]


def style_sheet(sheet, *, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in list(column)[:200]]
        width = min(max(max(map(len, values), default=8) + 2, 10), 42)
        sheet.column_dimensions[column[0].column_letter].width = width


def write_rows(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Reconcile material-matched keyword spend with account spend."
    )
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--date-from", type=date.fromisoformat, required=True)
    parser.add_argument("--date-to", type=date.fromisoformat, required=True)
    parser.add_argument("--task-id")
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    args = parser.parse_args()
    if args.date_to < args.date_from:
        raise ValueError("结束日期不能早于开始日期")

    params = {
        "project_code": args.project_code,
        "date_from": args.date_from,
        "date_to": args.date_to,
    }
    scope_cte = """
        WITH scoped_accounts AS (
            SELECT account.id, account.baidu_account_id, account.login_name,
                   account.account_subject
            FROM search_marketing.accounts AS account
            JOIN search_marketing.projects AS project
              ON project.id = account.project_id
            WHERE project.code = :project_code
              AND account.is_active
              AND account.account_type::text = 'SECOND_HOP'
        )
    """
    engine = create_engine(load_database_url(), pool_pre_ping=True)
    with engine.connect() as connection:
        accounts = fetch_rows(
            connection,
            scope_cte
            + """
            , account_spend AS (
                SELECT fact.account_id, sum(fact.spend) AS spend,
                       count(*) AS data_days,
                       min(fact.report_date) AS first_date,
                       max(fact.report_date) AS last_date
                FROM search_marketing.performance_daily AS fact
                JOIN scoped_accounts AS account ON account.id = fact.account_id
                WHERE fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.account_id
            ), keyword_spend AS (
                SELECT fact.account_id, sum(fact.spend) AS spend,
                       count(DISTINCT fact.report_date) AS data_days,
                       min(fact.report_date) AS first_date,
                       max(fact.report_date) AS last_date
                FROM search_marketing.keyword_performance_daily AS fact
                JOIN scoped_accounts AS account ON account.id = fact.account_id
                WHERE fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.account_id
            )
            SELECT account.account_subject, account.login_name,
                   account.baidu_account_id,
                   coalesce(account_spend.spend, 0) AS account_spend,
                   coalesce(keyword_spend.spend, 0) AS keyword_spend,
                   coalesce(account_spend.spend, 0)
                     - coalesce(keyword_spend.spend, 0) AS difference,
                   CASE WHEN coalesce(account_spend.spend, 0) = 0 THEN NULL
                        ELSE round(
                            coalesce(keyword_spend.spend, 0)
                            / account_spend.spend * 100, 4
                        ) END AS coverage_percent,
                   coalesce(account_spend.data_days, 0) AS account_data_days,
                   coalesce(keyword_spend.data_days, 0) AS keyword_data_days,
                   account_spend.first_date AS account_first_date,
                   account_spend.last_date AS account_last_date,
                   keyword_spend.first_date AS keyword_first_date,
                   keyword_spend.last_date AS keyword_last_date
            FROM scoped_accounts AS account
            LEFT JOIN account_spend ON account_spend.account_id = account.id
            LEFT JOIN keyword_spend ON keyword_spend.account_id = account.id
            ORDER BY abs(
                coalesce(account_spend.spend, 0)
                - coalesce(keyword_spend.spend, 0)
            ) DESC, account.baidu_account_id
            """,
            params,
        )
        daily = fetch_rows(
            connection,
            scope_cte
            + """
            , account_daily AS (
                SELECT fact.report_date, sum(fact.spend) AS spend,
                       count(DISTINCT fact.account_id) AS accounts
                FROM search_marketing.performance_daily AS fact
                JOIN scoped_accounts AS account ON account.id = fact.account_id
                WHERE fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.report_date
            ), keyword_daily AS (
                SELECT fact.report_date, sum(fact.spend) AS spend,
                       count(DISTINCT fact.account_id) AS accounts
                FROM search_marketing.keyword_performance_daily AS fact
                JOIN scoped_accounts AS account ON account.id = fact.account_id
                WHERE fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.report_date
            )
            SELECT day::date AS report_date,
                   coalesce(account_daily.spend, 0) AS account_spend,
                   coalesce(keyword_daily.spend, 0) AS keyword_spend,
                   coalesce(account_daily.spend, 0)
                     - coalesce(keyword_daily.spend, 0) AS difference,
                   CASE WHEN coalesce(account_daily.spend, 0) = 0 THEN NULL
                        ELSE round(
                            coalesce(keyword_daily.spend, 0)
                            / account_daily.spend * 100, 4
                        ) END AS coverage_percent,
                   coalesce(account_daily.accounts, 0) AS account_count,
                   coalesce(keyword_daily.accounts, 0) AS keyword_account_count
            FROM generate_series(
                CAST(:date_from AS date), CAST(:date_to AS date), interval '1 day'
            ) AS day
            LEFT JOIN account_daily ON account_daily.report_date = day::date
            LEFT JOIN keyword_daily ON keyword_daily.report_date = day::date
            ORDER BY day
            """,
            params,
        )
        account_day_differences = fetch_rows(
            connection,
            scope_cte
            + """
            , keyword_daily AS (
                SELECT fact.account_id, fact.report_date,
                       sum(fact.spend) AS spend
                FROM search_marketing.keyword_performance_daily AS fact
                JOIN scoped_accounts AS account ON account.id = fact.account_id
                WHERE fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.account_id, fact.report_date
            )
            SELECT account.account_subject, account.login_name,
                   account.baidu_account_id, fact.report_date,
                   fact.spend AS account_spend,
                   coalesce(keyword_daily.spend, 0) AS keyword_spend,
                   fact.spend - coalesce(keyword_daily.spend, 0) AS difference
            FROM search_marketing.performance_daily AS fact
            JOIN scoped_accounts AS account ON account.id = fact.account_id
            LEFT JOIN keyword_daily
              ON keyword_daily.account_id = fact.account_id
             AND keyword_daily.report_date = fact.report_date
            WHERE fact.report_date BETWEEN :date_from AND :date_to
              AND abs(fact.spend - coalesce(keyword_daily.spend, 0)) > 0.01
            ORDER BY abs(
                fact.spend - coalesce(keyword_daily.spend, 0)
            ) DESC, account.baidu_account_id, fact.report_date
            """,
            params,
        )
        keyword_only_days = fetch_rows(
            connection,
            scope_cte
            + """
            , keyword_daily AS (
                SELECT fact.account_id, fact.report_date,
                       sum(fact.spend) AS spend
                FROM search_marketing.keyword_performance_daily AS fact
                JOIN scoped_accounts AS account ON account.id = fact.account_id
                WHERE fact.report_date BETWEEN :date_from AND :date_to
                GROUP BY fact.account_id, fact.report_date
            )
            SELECT account.account_subject, account.login_name,
                   account.baidu_account_id, keyword_daily.report_date,
                   keyword_daily.spend AS keyword_spend
            FROM keyword_daily
            JOIN scoped_accounts AS account ON account.id = keyword_daily.account_id
            LEFT JOIN search_marketing.performance_daily AS fact
              ON fact.account_id = keyword_daily.account_id
             AND fact.report_date = keyword_daily.report_date
            WHERE fact.id IS NULL AND keyword_daily.spend <> 0
            ORDER BY keyword_daily.spend DESC
            """,
            params,
        )
        task_state: dict[str, Any] = {}
        if args.task_id:
            task_row = connection.execute(
                text(
                    """
                    SELECT status::text AS status, current_node, progress,
                           result, last_error
                    FROM search_marketing.background_tasks
                    WHERE id = :task_id
                    """
                ),
                {"task_id": args.task_id},
            ).mappings().first()
            task_state = dict(task_row or {})

    for row in accounts:
        row["account_spend"] = as_decimal(row["account_spend"])
        row["keyword_spend"] = as_decimal(row["keyword_spend"])
        row["difference"] = as_decimal(row["difference"])
    total_account = sum((row["account_spend"] for row in accounts), Decimal("0"))
    total_keyword = sum((row["keyword_spend"] for row in accounts), Decimal("0"))
    total_difference = total_account - total_keyword
    coverage = (
        (total_keyword / total_account * Decimal("100")).quantize(Decimal("0.0001"))
        if total_account
        else None
    )
    summary = {
        "project_code": args.project_code,
        "account_type": "二跳账户",
        "date_from": args.date_from,
        "date_to": args.date_to,
        "scope_accounts": len(accounts),
        "second_hop_account_spend": total_account,
        "keyword_spend": total_keyword,
        "difference_second_hop_account_minus_keyword": total_difference,
        "keyword_spend_coverage_percent": coverage,
        "accounts_exact_within_0_01": sum(
            abs(row["difference"]) <= MONEY_TOLERANCE for row in accounts
        ),
        "accounts_keyword_lower": sum(
            row["difference"] > MONEY_TOLERANCE for row in accounts
        ),
        "accounts_keyword_higher": sum(
            row["difference"] < -MONEY_TOLERANCE for row in accounts
        ),
        "account_day_differences": len(account_day_differences),
        "keyword_only_account_days": len(keyword_only_days),
        "task_status": task_state.get("status"),
        "task_node": task_state.get("current_node"),
        "task_state": (task_state.get("result") or {}).get("state", {}),
        "definition": (
            "关键词消耗仅包含去除[已删除]后与物料中心关键词精确匹配的数据；"
            "差额=二跳账户消耗-物料匹配关键词消耗。"
        ),
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "核对摘要"
    sheet.append(["项目", "值"])
    for key, value in summary.items():
        if key == "task_state":
            continue
        sheet.append([key, json_value(value)])
    style_sheet(sheet)

    account_rows = []
    for row in accounts:
        difference = row["difference"]
        result = (
            "一致"
            if abs(difference) <= MONEY_TOLERANCE
            else ("关键词消耗偏低" if difference > 0 else "关键词消耗偏高")
        )
        account_rows.append(
            [
                row["account_subject"],
                row["login_name"],
                row["baidu_account_id"],
                float(row["account_spend"]),
                float(row["keyword_spend"]),
                float(difference),
                row["coverage_percent"],
                row["account_data_days"],
                row["keyword_data_days"],
                row["account_first_date"],
                row["account_last_date"],
                row["keyword_first_date"],
                row["keyword_last_date"],
                result,
            ]
        )
    write_rows(
        workbook.create_sheet("按账户核对"),
        [
            "账号主体",
            "账户",
            "账户ID",
            "二跳账户消耗",
            "关键词消耗",
            "差额（二跳账户-关键词）",
            "关键词消耗覆盖率%",
            "账户数据天数",
            "关键词数据天数",
            "账户开始日期",
            "账户结束日期",
            "关键词开始日期",
            "关键词结束日期",
            "结论",
        ],
        account_rows,
    )
    write_rows(
        workbook.create_sheet("按日期核对"),
        [
            "日期",
            "二跳账户消耗",
            "关键词消耗",
            "差额（二跳账户-关键词）",
            "关键词消耗覆盖率%",
            "有账户数据账户数",
            "有关键词数据账户数",
        ],
        [list(row.values()) for row in daily],
    )
    write_rows(
        workbook.create_sheet("账户日期差异"),
        [
            "账号主体",
            "账户",
            "账户ID",
            "日期",
            "二跳账户消耗",
            "关键词消耗",
            "差额（二跳账户-关键词）",
        ],
        [list(row.values()) for row in account_day_differences],
    )
    write_rows(
        workbook.create_sheet("仅关键词有消耗"),
        ["账号主体", "账户", "账户ID", "日期", "关键词消耗"],
        [list(row.values()) for row in keyword_only_days],
    )

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_xlsx)
    args.output_json.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=json_value),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, default=json_value))


if __name__ == "__main__":
    main()
