from __future__ import annotations

import argparse
import html
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, text


SOURCE_LABELS = {
    "visitors": "实时访问详情",
    "copy_visitors": "复制实时详情",
    "conversions": "获客助手实时转化详情",
}
ACCOUNT_EXCLUSIONS = {
    "custom_id_not_found",
    "account_not_found",
    "remark_account_not_found",
    "remark_account_ambiguous",
    "missing_account_reference",
}
DELETED_MARKER = re.compile(r"(?:\[已删除\]|【已删除】)")
ZHANGHUID = re.compile(r"(?:^|[?&#])zhanghuid=([^?&#]*)", re.IGNORECASE)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="导出好多粉未匹配账户和关键词")
    parser.add_argument("--task-id", required=True)
    parser.add_argument("--project-code", default="weight-loss-search")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary-output", type=Path)
    return parser.parse_args()


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized_key(value: Any) -> str:
    return normalized_text(value).casefold()


def normalized_keyword(value: Any) -> str:
    text_value = DELETED_MARKER.sub("", normalized_text(value))
    return text_value.strip()


def extract_zhanghuid(value: Any) -> str:
    candidate = normalized_text(value)
    if not candidate:
        return ""
    for _ in range(3):
        candidate = html.unescape(candidate)
        parsed = urlparse(candidate)
        query = parse_qs(parsed.query, keep_blank_values=True)
        for key, values in query.items():
            if key.casefold() == "zhanghuid" and values:
                return normalized_text(re.split(r"[?&#]", values[0], maxsplit=1)[0])
        match = ZHANGHUID.search(candidate)
        if match:
            return normalized_text(unquote(match.group(1)))
        decoded = unquote(candidate)
        if decoded == candidate:
            break
        candidate = decoded
    return ""


def normalize_account_id(value: Any) -> str:
    result = normalized_key(value)
    if re.fullmatch(r"\d+\.0+", result):
        return result.split(".", 1)[0]
    return result


def gibberish_reason(value: str) -> str:
    if not value:
        return ""
    if "\ufffd" in value or "锟斤拷" in value or "ï¿½" in value:
        return "包含替换字符或典型乱码片段"
    if any(unicodedata.category(char) in {"Cc", "Cs"} for char in value):
        return "包含控制字符或无效Unicode字符"
    visible = [char for char in value if not char.isspace()]
    if visible:
        useful = sum(
            char.isalnum() or "\u3400" <= char <= "\u9fff"
            for char in visible
        )
        if useful == 0:
            return "全部为符号"
    if re.fullmatch(r"(?:%[0-9A-Fa-f]{2}){2,}", value):
        return "疑似未解码的URL编码"
    return ""


def load_raw_records(path: str) -> tuple[dict[str, dict[str, Any]], str]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    return (
        {str(row.get("id")): row for row in payload.get("records") or []},
        str(payload.get("source_file") or ""),
    )


def add_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    widths: dict[int, int] | None = None,
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in (widths or {}).items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def serialize(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def main() -> int:
    args = parse_args()
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL 未配置")

    engine = create_engine(database_url, pool_pre_ping=True)
    with engine.connect() as connection:
        project = connection.execute(
            text("SELECT id::text, name FROM search_marketing.projects WHERE code = :code"),
            {"code": args.project_code},
        ).mappings().one()
        project_id = project["id"]
        task = connection.execute(
            text(
                """
                SELECT id::text, status::text, created_at, result
                FROM search_marketing.background_tasks
                WHERE id = CAST(:task_id AS uuid)
                """
            ),
            {"task_id": args.task_id},
        ).mappings().one()
        accounts = connection.execute(
            text(
                """
                SELECT id::text, baidu_account_id::text, login_name,
                       account_type::text,
                       coalesce(account_subject, '') AS account_subject
                FROM search_marketing.accounts
                WHERE project_id = CAST(:project_id AS uuid)
                ORDER BY login_name
                """
            ),
            {"project_id": project_id},
        ).mappings().all()
        custom_mappings = connection.execute(
            text(
                """
                SELECT mapping.custom_id, account.baidu_account_id::text,
                       account.login_name, account.account_subject
                FROM search_marketing.hduofen_account_mappings AS mapping
                JOIN search_marketing.accounts AS account ON account.id = mapping.account_id
                WHERE mapping.project_id = CAST(:project_id AS uuid)
                """
            ),
            {"project_id": project_id},
        ).mappings().all()
        materials = connection.execute(
            text(
                """
                SELECT id::text, keyword_text, campaign_name
                FROM search_marketing.material_keywords
                WHERE project_id = CAST(:project_id AS uuid)
                """
            ),
            {"project_id": project_id},
        ).mappings().all()

        account_by_numeric = {
            normalize_account_id(row["baidu_account_id"]): row for row in accounts
        }
        account_by_custom = {
            normalize_account_id(row["custom_id"]): row for row in custom_mappings
        }
        accounts_by_login: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in accounts:
            accounts_by_login[normalized_key(row["login_name"])].append(row)
        material_by_normalized: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in materials:
            material_by_normalized[normalized_keyword(row["keyword_text"])].append(row)

        statement = text(
            """
            SELECT event.source_type, event.source_event_id, event.event_date,
                   event.account_id::text, event.tracking_keyword,
                   event.material_keyword_id::text, event.exclusion_reason,
                   event.raw_data_path, account.login_name AS matched_account_name,
                   account.baidu_account_id::text AS matched_baidu_account_id
            FROM search_marketing.hduofen_events AS event
            LEFT JOIN search_marketing.accounts AS account ON account.id = event.account_id
            WHERE event.capture_task_id = CAST(:task_id AS uuid)
              AND (
                event.account_id IS NULL
                OR nullif(btrim(event.tracking_keyword), '') IS NOT NULL
              )
            ORDER BY event.raw_data_path, event.source_event_id
            """
        )
        result = connection.execution_options(stream_results=True).execute(
            statement, {"task_id": args.task_id}
        ).mappings()

        account_details: list[list[Any]] = []
        blank_account_details: list[list[Any]] = []
        keyword_details: list[list[Any]] = []
        normalized_keyword_matches: list[list[Any]] = []
        account_aggregates: dict[tuple[str, str, str], dict[str, Any]] = {}
        keyword_aggregates: dict[str, dict[str, Any]] = {}
        blank_keyword_counts: Counter[tuple[str, str]] = Counter()
        current_path = ""
        raw_records: dict[str, dict[str, Any]] = {}
        source_file = ""
        raw_missing = 0
        account_suggestion_cache: dict[tuple[str, str], tuple[str, str, str]] = {}

        def suggest_account(zhanghuid: str, remark: str) -> tuple[str, str, str]:
            cache_key = (normalize_account_id(zhanghuid), normalized_key(remark))
            if cache_key in account_suggestion_cache:
                return account_suggestion_cache[cache_key]
            if zhanghuid:
                key = normalize_account_id(zhanghuid)
                candidate = account_by_custom.get(key) or account_by_numeric.get(key)
                if candidate:
                    suggestion = (
                        str(candidate["login_name"]),
                        str(candidate.get("baidu_account_id") or ""),
                        "ID归一化后匹配",
                    )
                    account_suggestion_cache[cache_key] = suggestion
                    return suggestion
            if remark:
                candidates = accounts_by_login.get(normalized_key(remark)) or []
                if len(candidates) == 1:
                    candidate = candidates[0]
                    suggestion = (
                        str(candidate["login_name"]),
                        str(candidate["baidu_account_id"]),
                        "备注归一化后匹配",
                    )
                    account_suggestion_cache[cache_key] = suggestion
                    return suggestion
                remark_key = normalized_key(remark)
                prefix_candidates = [
                    account
                    for account in accounts
                    if normalized_key(account["login_name"]).startswith(remark_key)
                ]
                if remark_key.startswith("k基木鱼-"):
                    prefix_candidates = [
                        account
                        for account in prefix_candidates
                        if account["account_type"] == "SECOND_HOP"
                    ]
                if len(remark_key) >= 8 and len(prefix_candidates) == 1:
                    candidate = prefix_candidates[0]
                    suggestion = (
                        str(candidate["login_name"]),
                        str(candidate["baidu_account_id"]),
                        "好多粉备注截断，唯一账户名前缀匹配",
                    )
                    account_suggestion_cache[cache_key] = suggestion
                    return suggestion
                if len(remark_key) >= 8 and prefix_candidates:
                    suggestion = (
                        "；".join(
                            sorted({str(row["login_name"]) for row in prefix_candidates})[:10]
                        ),
                        "",
                        "账户名前缀存在多个候选，禁止自动匹配",
                    )
                    account_suggestion_cache[cache_key] = suggestion
                    return suggestion
            suggestion = ("", "", "")
            account_suggestion_cache[cache_key] = suggestion
            return suggestion

        for event in result:
            path = str(event["raw_data_path"])
            if path != current_path:
                current_path = path
                raw_records, source_file = load_raw_records(path)
            raw = raw_records.get(str(event["source_event_id"]))
            if raw is None:
                raw_missing += 1
                raw = {}

            complete_url = normalized_text(raw.get("complete_url") or raw.get("url"))
            zhanghuid = extract_zhanghuid(complete_url)
            remark = normalized_text(raw.get("account_remark"))
            keyword = normalized_text(event.get("tracking_keyword"))
            normalized_kw = normalized_keyword(keyword)
            file_name = normalized_text(raw.get("history_file")) or Path(source_file).name
            row_number = raw.get("history_row")
            source_label = SOURCE_LABELS.get(str(event["source_type"]), str(event["source_type"]))

            if event["account_id"] is None and event["exclusion_reason"] in ACCOUNT_EXCLUSIONS:
                suggestion, suggested_id, suggestion_reason = suggest_account(zhanghuid, remark)
                row = [
                    source_label,
                    serialize(event["event_date"]),
                    file_name,
                    row_number,
                    zhanghuid,
                    remark,
                    keyword,
                    complete_url,
                    event["exclusion_reason"],
                    suggestion,
                    suggested_id,
                    suggestion_reason,
                ]
                if not zhanghuid and not remark:
                    blank_account_details.append(row)
                else:
                    account_details.append(row)
                    account_type = "zhanghuid" if zhanghuid else "备注/url备注"
                    reference = zhanghuid or remark
                    aggregate_key = (account_type, reference, remark)
                    aggregate = account_aggregates.setdefault(
                        aggregate_key,
                        {
                            "count": 0,
                            "sources": set(),
                            "first_date": event["event_date"],
                            "last_date": event["event_date"],
                            "keywords": set(),
                            "sample_file": file_name,
                            "sample_row": row_number,
                            "sample_url": complete_url,
                            "reason": event["exclusion_reason"],
                            "suggestion": suggestion,
                            "suggested_id": suggested_id,
                            "suggestion_reason": suggestion_reason,
                        },
                    )
                    aggregate["count"] += 1
                    aggregate["sources"].add(source_label)
                    aggregate["first_date"] = min(aggregate["first_date"], event["event_date"])
                    aggregate["last_date"] = max(aggregate["last_date"], event["event_date"])
                    if keyword:
                        aggregate["keywords"].add(keyword)

            if not keyword:
                account_state = "账户已匹配" if event["account_id"] else "账户未匹配"
                blank_keyword_counts[(source_label, account_state)] += 1
                continue

            material_candidates = material_by_normalized.get(normalized_kw) or []
            exact_material = next(
                (row for row in material_candidates if row["keyword_text"] == keyword), None
            )
            if exact_material is not None:
                continue
            if material_candidates:
                normalized_keyword_matches.append(
                    [
                        keyword,
                        normalized_kw,
                        "；".join(sorted({str(row["keyword_text"]) for row in material_candidates})),
                        source_label,
                        event["matched_account_name"] or "",
                        file_name,
                        row_number,
                    ]
                )
                continue

            gibberish = gibberish_reason(keyword)
            account_reference = (
                str(event["matched_account_name"] or "") or zhanghuid or remark
            )
            detail = [
                keyword,
                normalized_kw,
                gibberish or "非乱码，需补充或修正",
                source_label,
                serialize(event["event_date"]),
                event["matched_account_name"] or "",
                event["matched_baidu_account_id"] or "",
                zhanghuid,
                remark,
                file_name,
                row_number,
                complete_url,
                event["exclusion_reason"] or "",
            ]
            keyword_details.append(detail)
            aggregate = keyword_aggregates.setdefault(
                keyword,
                {
                    "normalized": normalized_kw,
                    "gibberish": gibberish,
                    "count": 0,
                    "sources": set(),
                    "accounts": set(),
                    "matched_account_records": 0,
                    "first_date": event["event_date"],
                    "last_date": event["event_date"],
                    "sample_file": file_name,
                    "sample_row": row_number,
                    "sample_url": complete_url,
                },
            )
            aggregate["count"] += 1
            aggregate["sources"].add(source_label)
            if account_reference:
                aggregate["accounts"].add(account_reference)
            if event["account_id"]:
                aggregate["matched_account_records"] += 1
            aggregate["first_date"] = min(aggregate["first_date"], event["event_date"])
            aggregate["last_date"] = max(aggregate["last_date"], event["event_date"])

    account_summary_rows: list[list[Any]] = []
    for (reference_type, reference, remark), aggregate in sorted(
        account_aggregates.items(), key=lambda item: (-item[1]["count"], item[0])
    ):
        account_summary_rows.append(
            [
                reference_type,
                reference,
                remark,
                aggregate["count"],
                "、".join(sorted(aggregate["sources"])),
                serialize(aggregate["first_date"]),
                serialize(aggregate["last_date"]),
                len(aggregate["keywords"]),
                aggregate["reason"],
                aggregate["suggestion"],
                aggregate["suggested_id"],
                aggregate["suggestion_reason"],
                aggregate["sample_file"],
                aggregate["sample_row"],
                aggregate["sample_url"],
            ]
        )

    keyword_summary_rows: list[list[Any]] = []
    gibberish_summary_rows: list[list[Any]] = []
    for keyword, aggregate in sorted(
        keyword_aggregates.items(), key=lambda item: (-item[1]["count"], item[0])
    ):
        row = [
            keyword,
            aggregate["normalized"],
            aggregate["gibberish"] or "非乱码，需补充或修正",
            aggregate["count"],
            aggregate["matched_account_records"],
            len(aggregate["accounts"]),
            "、".join(sorted(aggregate["sources"])),
            serialize(aggregate["first_date"]),
            serialize(aggregate["last_date"]),
            "；".join(sorted(aggregate["accounts"])[:20]),
            aggregate["sample_file"],
            aggregate["sample_row"],
            aggregate["sample_url"],
        ]
        if aggregate["gibberish"]:
            gibberish_summary_rows.append(row)
        else:
            keyword_summary_rows.append(row)

    workbook = Workbook()
    workbook.remove(workbook.active)
    overview = [
        ["项目", project["name"]],
        ["导入任务", args.task_id],
        ["任务状态", task["status"]],
        ["未匹配账户记录（账户字段非空）", len(account_details)],
        ["未匹配账户引用组合", len(account_summary_rows)],
        ["账户字段为空记录", len(blank_account_details)],
        ["未匹配非乱码关键词记录", sum(row[3] for row in keyword_summary_rows)],
        ["未匹配非乱码关键词数", len(keyword_summary_rows)],
        ["疑似乱码关键词记录", sum(row[3] for row in gibberish_summary_rows)],
        ["疑似乱码关键词数", len(gibberish_summary_rows)],
        ["归一化后可匹配关键词记录", len(normalized_keyword_matches)],
        ["原始记录缺失", raw_missing],
        ["账户匹配规则", "zhanghuid优先匹配自定义ID/百度账户ID；无zhanghuid时匹配备注或url备注"],
        ["关键词匹配规则", "去除首尾空格、统一全半角并去掉[已删除]/【已删除】后，与当前物料中心精确匹配"],
    ]
    add_sheet(workbook, "核对说明", ["项目", "结果"], overview, {1: 32, 2: 95})
    add_sheet(
        workbook,
        "未匹配账户汇总",
        ["引用类型", "账户引用", "备注/url备注", "记录数", "来源表", "最早日期", "最晚日期", "关键词数", "导入排除原因", "建议账户", "建议账户ID", "建议依据", "样本文件", "样本行", "样本URL"],
        account_summary_rows,
        {1: 16, 2: 28, 3: 28, 4: 12, 5: 28, 10: 32, 13: 42, 15: 80},
    )
    add_sheet(
        workbook,
        "未匹配账户明细",
        ["来源表", "日期", "文件", "行号", "zhanghuid", "备注/url备注", "追踪关键词", "完整URL", "导入排除原因", "建议账户", "建议账户ID", "建议依据"],
        account_details,
        {1: 24, 3: 42, 5: 24, 6: 28, 7: 30, 8: 90, 10: 32, 12: 28},
    )
    add_sheet(
        workbook,
        "账户字段为空",
        ["来源表", "日期", "文件", "行号", "zhanghuid", "备注/url备注", "追踪关键词", "完整URL", "导入排除原因", "建议账户", "建议账户ID", "建议依据"],
        blank_account_details,
        {1: 24, 3: 42, 7: 30, 8: 90},
    )
    keyword_headers = ["追踪关键词", "归一化关键词", "判定", "记录数", "账户已匹配记录数", "涉及账户数", "来源表", "最早日期", "最晚日期", "涉及账户（最多20个）", "样本文件", "样本行", "样本URL"]
    add_sheet(workbook, "未匹配关键词汇总", keyword_headers, keyword_summary_rows, {1: 38, 2: 38, 3: 28, 7: 28, 10: 70, 11: 42, 13: 90})
    add_sheet(workbook, "疑似乱码关键词", keyword_headers, gibberish_summary_rows, {1: 38, 2: 38, 3: 32, 7: 28, 10: 70, 11: 42, 13: 90})
    add_sheet(
        workbook,
        "未匹配关键词明细",
        ["追踪关键词", "归一化关键词", "判定", "来源表", "日期", "已匹配账户", "百度账户ID", "zhanghuid", "备注/url备注", "文件", "行号", "完整URL", "导入排除原因"],
        keyword_details,
        {1: 38, 2: 38, 3: 28, 4: 24, 6: 30, 8: 24, 9: 28, 10: 42, 12: 90},
    )
    add_sheet(
        workbook,
        "归一化后已匹配关键词",
        ["原追踪关键词", "归一化关键词", "物料中心关键词", "来源表", "已匹配账户", "文件", "行号"],
        normalized_keyword_matches,
        {1: 38, 2: 38, 3: 38, 4: 24, 5: 30, 6: 42},
    )
    blank_rows = [
        [source, state, count]
        for (source, state), count in sorted(blank_keyword_counts.items())
    ]
    add_sheet(workbook, "关键词为空统计", ["来源表", "账户状态", "记录数"], blank_rows, {1: 28, 2: 18, 3: 14})

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    # Reopen the workbook so a corrupt or incomplete XLSX fails the command.
    check = load_workbook(args.output, read_only=True, data_only=True)
    sheet_rows = {sheet.title: sheet.max_row - 1 for sheet in check.worksheets}
    check.close()

    summary = {
        "project": project["name"],
        "task_id": args.task_id,
        "task_status": str(task["status"]),
        "unmatched_account_records_nonblank": len(account_details),
        "unmatched_account_reference_groups": len(account_summary_rows),
        "blank_account_records": len(blank_account_details),
        "unmatched_keyword_records_non_gibberish": sum(row[3] for row in keyword_summary_rows),
        "unmatched_keywords_non_gibberish": len(keyword_summary_rows),
        "suspected_gibberish_keyword_records": sum(row[3] for row in gibberish_summary_rows),
        "suspected_gibberish_keywords": len(gibberish_summary_rows),
        "normalized_keyword_match_records": len(normalized_keyword_matches),
        "raw_records_missing": raw_missing,
        "sheet_rows": sheet_rows,
        "output": str(args.output),
    }
    summary_output = args.summary_output or args.output.with_suffix(".json")
    summary_output.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
