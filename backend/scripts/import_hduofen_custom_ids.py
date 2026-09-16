"""Import project-scoped Hduofen custom IDs by exact account login name."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import uuid
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook


BACKEND_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = BACKEND_ROOT.parent
sys.path.insert(0, str(BACKEND_ROOT / "src"))


def configure_local_database() -> None:
    if os.getenv("DATABASE_URL"):
        return
    for line in (PROJECT_ROOT / ".env").read_text(encoding="utf-8").splitlines():
        if line.startswith("DATABASE_URL="):
            os.environ["DATABASE_URL"] = line.split("=", 1)[1].strip().replace(
                "@postgres:", "@127.0.0.1:"
            )
            return


def clean(value) -> str:
    return str(value or "").strip()


def read_mapping_rows(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows, ())]
    if "自定义ID" not in headers or "账户" not in headers:
        workbook.close()
        raise ValueError("自定义ID工作簿缺少“自定义ID”或“账户”列")
    id_index = headers.index("自定义ID")
    account_index = headers.index("账户")
    result = []
    for values in rows:
        custom_id = clean(values[id_index] if id_index < len(values) else None)
        account_name = clean(values[account_index] if account_index < len(values) else None)
        if custom_id or account_name:
            result.append((custom_id, account_name))
    workbook.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--project-code", default="jfsem")
    parser.add_argument("--report-path", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    configure_local_database()

    from sqlalchemy import select
    from sqlalchemy.dialects.postgresql import insert as pg_insert

    from search_console.db import SessionLocal
    from search_console.models import (
        Account,
        AuditEvent,
        HduofenAccountMapping,
        Project,
    )

    path = Path(args.file)
    rows = read_mapping_rows(path)
    source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    id_counts = Counter(custom_id for custom_id, _ in rows)
    account_counts = Counter(account_name for _, account_name in rows)
    duplicate_ids = sorted(key for key, count in id_counts.items() if key and count > 1)
    duplicate_accounts = sorted(key for key, count in account_counts.items() if key and count > 1)
    blank_rows = sum(not custom_id or not account_name for custom_id, account_name in rows)

    with SessionLocal() as db:
        project = db.scalar(select(Project).where(Project.code == args.project_code))
        if project is None:
            raise SystemExit(f"项目不存在：{args.project_code}")
        accounts = db.scalars(select(Account).where(Account.project_id == project.id)).all()
        by_name: dict[str, list[Account]] = {}
        for account in accounts:
            by_name.setdefault(account.login_name.strip(), []).append(account)

        matched = []
        unmatched = []
        ambiguous = []
        for custom_id, account_name in rows:
            candidates = by_name.get(account_name) or []
            if not custom_id or not account_name or not candidates:
                unmatched.append({"custom_id": custom_id, "account": account_name})
            elif len(candidates) > 1:
                ambiguous.append({"custom_id": custom_id, "account": account_name})
            else:
                matched.append((custom_id, candidates[0]))

        existing = {
            row.custom_id: row
            for row in db.scalars(select(HduofenAccountMapping).where(
                HduofenAccountMapping.project_id == project.id
            )).all()
        }
        reassigned = sum(
            custom_id in existing and existing[custom_id].account_id != account.id
            for custom_id, account in matched
        )
        profile = {
            "generated_at": datetime.now(UTC).isoformat(),
            "source_file": str(path),
            "source_sha256": source_sha256,
            "rows": len(rows),
            "matched": len(matched),
            "unmatched": unmatched,
            "ambiguous": ambiguous,
            "blank_rows": blank_rows,
            "duplicate_ids": duplicate_ids,
            "duplicate_accounts": duplicate_accounts,
            "existing_mappings": len(existing),
            "reassigned_mappings": reassigned,
            "apply_requested": args.apply,
        }
        report_path = Path(args.report_path)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(profile, ensure_ascii=False))

        if not args.apply:
            return
        if blank_rows or duplicate_ids or duplicate_accounts or unmatched or ambiguous:
            raise SystemExit("自定义ID预检未通过，已拒绝写库")

        now = datetime.now(UTC)
        values = [{
            "id": uuid.uuid4(),
            "project_id": project.id,
            "custom_id": custom_id,
            "account_id": account.id,
            "source_file": str(path),
            "source_sha256": source_sha256,
            "created_at": now,
            "updated_at": now,
        } for custom_id, account in matched]
        statement = pg_insert(HduofenAccountMapping).values(values)
        excluded = statement.excluded
        db.execute(statement.on_conflict_do_update(
            index_elements=[
                HduofenAccountMapping.project_id,
                HduofenAccountMapping.custom_id,
            ],
            set_={
                "account_id": excluded.account_id,
                "source_file": excluded.source_file,
                "source_sha256": excluded.source_sha256,
                "updated_at": excluded.updated_at,
            },
        ))
        db.add(AuditEvent(
            project_id=project.id,
            actor="hduofen-custom-id-importer",
            action="hduofen.custom_id.import",
            target_type="project",
            target_id=str(project.id),
            summary="导入好多粉自定义ID账户映射",
            details={
                "source_file": str(path),
                "source_sha256": source_sha256,
                "mappings": len(values),
                "reassigned": reassigned,
            },
        ))
        db.commit()
        print(json.dumps({"imported": len(values), "reassigned": reassigned}, ensure_ascii=False))


if __name__ == "__main__":
    main()
