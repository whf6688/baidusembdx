from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


SHANGHAI = ZoneInfo("Asia/Shanghai")
RANGE_FILE_PATTERN = re.compile(
    r"^(?P<label>实时访问详情|复制实时详情)"
    r"(?P<start>\d{6}_\d{2}_\d{2}_\d{2})至(?P<end>\d{6}_\d{2}_\d{2}_\d{2})\.xls$"
)
CONVERSION_FILE_PATTERN = re.compile(
    r"^获客助手实时转化详情(?P<start>\d{4}-\d{2}-\d{2})至(?P<end>\d{4}-\d{2}-\d{2})"
    r"(?: \(\d+\))?\.xls$"
)
SOURCE_BY_LABEL = {
    "实时访问详情": "visitors",
    "复制实时详情": "copy_visitors",
}
SOURCE_COLUMNS = {
    "visitors": {
        "url": "访问链接",
        "remark": "落地页备注",
        "keyword": "追踪关键词",
        "event_time": "访问时间",
    },
    "copy_visitors": {
        "url": "访问链接",
        "remark": "备注",
        "keyword": "追踪关键词",
        "event_time": "复制时间",
    },
    "conversions": {
        "url": "访问url",
        "remark": "url备注",
        "keyword": "追踪关键词",
        "event_time": "到粉时间",
        "fallback_times": ("申请添加时间", "访问时间"),
    },
}


@dataclass(frozen=True)
class HistoryFile:
    path: Path
    source_type: str
    start_at: datetime
    end_at: datetime


@dataclass(frozen=True)
class HistoryDataset:
    history_file: HistoryFile
    file_sha256: str
    headers: tuple[str, ...]
    records: tuple[dict[str, Any], ...]
    blank_rows: int
    outside_filename_window: int


def clean_cell(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _local_datetime(value: Any, fallback: datetime) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, str) and value.strip():
        try:
            parsed = datetime.fromisoformat(value.strip().replace("/", "-"))
        except ValueError:
            return fallback
    else:
        return fallback
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=SHANGHAI)
    return parsed.astimezone(SHANGHAI)


def parse_history_filename(path: Path) -> HistoryFile:
    if match := RANGE_FILE_PATTERN.fullmatch(path.name):
        start_at = datetime.strptime(match.group("start"), "%y%m%d_%H_%M_%S").replace(
            tzinfo=SHANGHAI
        )
        end_at = datetime.strptime(match.group("end"), "%y%m%d_%H_%M_%S").replace(
            tzinfo=SHANGHAI
        )
        source_type = SOURCE_BY_LABEL[match.group("label")]
    elif match := CONVERSION_FILE_PATTERN.fullmatch(path.name):
        start_day = date.fromisoformat(match.group("start"))
        end_day = date.fromisoformat(match.group("end"))
        start_at = datetime.combine(start_day, time.min, tzinfo=SHANGHAI)
        end_at = datetime.combine(end_day, time(23, 59, 59), tzinfo=SHANGHAI)
        source_type = "conversions"
    else:
        raise ValueError(f"无法识别好多粉历史文件名：{path.name}")
    if end_at < start_at:
        raise ValueError(f"好多粉历史文件时间范围倒置：{path.name}")
    return HistoryFile(path=path, source_type=source_type, start_at=start_at, end_at=end_at)


def discover_history_files(root: Path) -> list[HistoryFile]:
    if not root.is_dir():
        raise FileNotFoundError(f"好多粉历史目录不存在：{root}")
    files = [parse_history_filename(path) for path in root.rglob("*.xls") if path.is_file()]
    return sorted(files, key=lambda item: (item.start_at, item.source_type, item.path.name))


def read_history_workbook(history_file: HistoryFile) -> HistoryDataset:
    file_sha256 = hashlib.sha256(history_file.path.read_bytes()).hexdigest()
    with history_file.path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=False, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = tuple(clean_cell(value) or "" for value in next(rows, ()))
    positions = {header: index for index, header in enumerate(headers) if header}
    spec = SOURCE_COLUMNS[history_file.source_type]
    required = {str(spec["url"]), str(spec["remark"]), str(spec["keyword"]), str(spec["event_time"])}
    missing = sorted(required - positions.keys())
    if missing:
        workbook.close()
        raise ValueError(f"{history_file.path.name} 缺少字段：{missing}")

    records: list[dict[str, Any]] = []
    blank_rows = 0
    outside_filename_window = 0
    fallback_times = tuple(spec.get("fallback_times") or ())
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            blank_rows += 1
            continue

        def value_for(header: str) -> Any:
            index = positions.get(header)
            return values[index] if index is not None and index < len(values) else None

        raw_time = value_for(str(spec["event_time"]))
        if raw_time in (None, ""):
            for header in fallback_times:
                raw_time = value_for(header)
                if raw_time not in (None, ""):
                    break
        event_at = _local_datetime(raw_time, history_file.start_at)
        if event_at < history_file.start_at or event_at > history_file.end_at:
            outside_filename_window += 1
        record_id = hashlib.sha256(
            f"{file_sha256}:{history_file.source_type}:{row_number}".encode("utf-8")
        ).hexdigest()
        record = {
            "id": record_id,
            "complete_url": clean_cell(value_for(str(spec["url"]))),
            "account_remark": clean_cell(value_for(str(spec["remark"]))),
            "keyword_encoded": clean_cell(value_for(str(spec["keyword"]))),
            "start_time": event_at.isoformat(),
            "history_file": history_file.path.name,
            "history_row": row_number,
            "history_exclusion_reason": (
                "outside_filename_window"
                if event_at < history_file.start_at or event_at > history_file.end_at
                else None
            ),
        }
        if history_file.source_type == "conversions":
            record["isAdd_time"] = event_at.isoformat()
        records.append(record)
    workbook.close()
    return HistoryDataset(
        history_file=history_file,
        file_sha256=file_sha256,
        headers=headers,
        records=tuple(records),
        blank_rows=blank_rows,
        outside_filename_window=outside_filename_window,
    )


def write_history_dataset(dataset: HistoryDataset, storage_root: Path) -> Path:
    relative_dir = (
        storage_root
        / "hduofen_history"
        / dataset.history_file.source_type
        / dataset.history_file.start_at.strftime("%Y/%m")
    )
    relative_dir.mkdir(parents=True, exist_ok=True)
    output = relative_dir / f"{dataset.history_file.path.stem}.json"
    payload = {
        "source_file": str(dataset.history_file.path),
        "source_sha256": dataset.file_sha256,
        "captured_at": dataset.history_file.end_at.astimezone(UTC).isoformat(),
        "date_from": dataset.history_file.start_at.isoformat(),
        "date_to": dataset.history_file.end_at.isoformat(),
        "expected_count": len(dataset.records),
        "record_count": len(dataset.records),
        "complete": True,
        "headers": list(dataset.headers),
        "records": list(dataset.records),
    }
    output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return output
