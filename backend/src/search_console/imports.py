import hashlib
import io
from collections import defaultdict
from typing import Any

from openpyxl import load_workbook

from .keyword_encoding import KEYWORD_ENCODING_VERSION, encode_keyword_utf8
from .keyword_tiers import normalize_tier_name
from .schemas import ImportPreview

MAX_KEYWORDS_PER_UNIT = 5000
REQUIRED_HEADERS = {"计划", "关键词"}


def parse_material_workbook(
    filename: str,
    content: bytes,
    target_accounts: list[dict[str, Any]] | None = None,
) -> tuple[ImportPreview, list[dict[str, Any]]]:
    """Parse and validate the workbook in one streaming pass.

    The returned normalized rows are the only input used for PostgreSQL writes;
    later business processing reads the database and never reopens the workbook.
    """
    digest = hashlib.sha256(content).hexdigest()
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        preview = ImportPreview(
            filename=filename, sha256=digest, row_count=0, valid_count=0, invalid_count=1,
            create_count=0, update_count=0, skip_count=0,
            errors=[{"row": 1, "message": f"缺少列：{', '.join(sorted(missing))}"}], unit_chunks=[]
        )
        return preview, []

    positions = {name: headers.index(name) for name in REQUIRED_HEADERS}
    errors: list[dict] = []
    seen: set[tuple[str, str]] = set()
    grouped: dict[str, int] = defaultdict(int)
    normalized_rows: list[dict[str, Any]] = []
    valid = skipped = total = 0
    for row_number, row in enumerate(rows, start=2):
        total += 1
        campaign = normalize_tier_name(str(row[positions["计划"]] or "").strip())
        keyword = str(row[positions["关键词"]] or "").strip()
        row_errors = []
        if not campaign or not keyword:
            row_errors.append("计划和关键词不能为空")
        unique_key = (campaign.casefold(), keyword.casefold())
        if unique_key in seen:
            skipped += 1
            continue
        seen.add(unique_key)
        if row_errors:
            errors.append({"row": row_number, "message": "；".join(row_errors)})
        else:
            valid += 1
            grouped[campaign] += 1
            normalized_rows.append({
                "row_number": row_number,
                "campaign_name": campaign,
                "keyword_text": keyword,
                "keyword_utf8_encoded": encode_keyword_utf8(keyword),
                "keyword_encoding_version": KEYWORD_ENCODING_VERSION,
            })

    chunks = []
    accounts = target_accounts or []
    if accounts:
        for account in accounts:
            for campaign, count in grouped.items():
                chunk_count = (count + MAX_KEYWORDS_PER_UNIT - 1) // MAX_KEYWORDS_PER_UNIT
                chunks.append({
                    "account_id": account["account_id"],
                    "login_name": account["login_name"],
                    "landing_url_template": account["landing_url_template"],
                    "campaign": campaign,
                    "keyword_count": count,
                    "unit_count": chunk_count,
                })
    else:
        for campaign, count in grouped.items():
            chunk_count = (count + MAX_KEYWORDS_PER_UNIT - 1) // MAX_KEYWORDS_PER_UNIT
            chunks.append({
                "account_id": None,
                "campaign": campaign,
                "keyword_count": count,
                "unit_count": chunk_count,
            })
    target_count = len(accounts)
    preview = ImportPreview(
        filename=filename, sha256=digest, row_count=total, valid_count=valid,
        invalid_count=len(errors), create_count=valid * (target_count or 1),
        update_count=0, skip_count=skipped,
        errors=errors[:100], unit_chunks=chunks, target_accounts=accounts,
    )
    return preview, normalized_rows


def filter_new_material_rows(
    preview: ImportPreview,
    rows: list[dict[str, Any]],
    existing_keywords: set[str],
) -> tuple[ImportPreview, list[dict[str, Any]]]:
    """Keep only rows that are new to the project-wide material library."""
    normalized_existing = {keyword.casefold() for keyword in existing_keywords}
    new_rows: list[dict[str, Any]] = []
    grouped: dict[str, int] = defaultdict(int)
    skipped = preview.skip_count
    for row in rows:
        key = row["keyword_text"].casefold()
        if key in normalized_existing:
            skipped += 1
            continue
        normalized_existing.add(key)
        new_rows.append(row)
        grouped[row["campaign_name"]] += 1

    unit_chunks = [
        {
            "account_id": None,
            "campaign": campaign,
            "keyword_count": count,
            "unit_count": (count + MAX_KEYWORDS_PER_UNIT - 1) // MAX_KEYWORDS_PER_UNIT,
        }
        for campaign, count in grouped.items()
    ]
    updated = preview.model_copy(update={
        "valid_count": len(new_rows),
        "create_count": len(new_rows),
        "skip_count": skipped,
        "unit_chunks": unit_chunks,
    })
    return updated, new_rows


def preview_workbook(
    filename: str,
    content: bytes,
    target_accounts: list[dict[str, Any]] | None = None,
) -> ImportPreview:
    preview, _ = parse_material_workbook(filename, content, target_accounts)
    return preview
