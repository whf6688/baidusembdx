import hashlib
import io
import unicodedata
import uuid
from collections import defaultdict
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation, ROUND_FLOOR, ROUND_HALF_UP
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, insert, select
from sqlalchemy.orm import Session

from .models import AuditEvent, MaterialKeyword, MaterialKeywordPerformanceYearly


HISTORICAL_HEADERS = ("关键词", "展现", "点击", "消费", "UV", "复制", "加粉")
COUNT_FIELDS = ("impressions", "clicks", "uv", "copies", "adds")


def normalize_keyword_key(value: Any) -> str:
    return unicodedata.normalize("NFKC", str(value or "")).strip().casefold()


def decimal_value(value: Any, *, sheet: str, row_number: int, field: str) -> Decimal:
    try:
        parsed = Decimal(str(value if value not in (None, "") else 0))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{sheet}第{row_number}行{field}不是有效数字") from exc
    if not parsed.is_finite() or parsed < 0:
        raise ValueError(f"{sheet}第{row_number}行{field}不能为负数或非有限值")
    return parsed


def floor_count(value: Any, *, sheet: str, row_number: int, field: str) -> tuple[int, bool]:
    parsed = decimal_value(value, sheet=sheet, row_number=row_number, field=field)
    floored = parsed.to_integral_value(rounding=ROUND_FLOOR)
    return int(floored), parsed != floored


def parse_historical_keyword_performance(content: bytes) -> tuple[dict[str, dict], dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    aggregated: dict[str, dict] = {}
    input_rows = 0
    fractional_counts: dict[str, int] = defaultdict(int)
    sheet_rows: dict[str, int] = {}

    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if headers[: len(HISTORICAL_HEADERS)] != HISTORICAL_HEADERS:
            raise ValueError(
                f"{worksheet.title}表头必须依次为：{'、'.join(HISTORICAL_HEADERS)}"
            )
        current_sheet_rows = 0
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row[:7]):
                continue
            keyword_text = str(row[0] or "").strip()
            keyword_key = normalize_keyword_key(keyword_text)
            if not keyword_key:
                raise ValueError(f"{worksheet.title}第{row_number}行关键词为空")

            impressions, impressions_fractional = floor_count(
                row[1], sheet=worksheet.title, row_number=row_number, field="展现"
            )
            clicks, clicks_fractional = floor_count(
                row[2], sheet=worksheet.title, row_number=row_number, field="点击"
            )
            spend = decimal_value(
                row[3], sheet=worksheet.title, row_number=row_number, field="消费"
            )
            uv, uv_fractional = floor_count(
                row[4], sheet=worksheet.title, row_number=row_number, field="UV"
            )
            copies, copies_fractional = floor_count(
                row[5], sheet=worksheet.title, row_number=row_number, field="复制"
            )
            adds, adds_fractional = floor_count(
                row[6], sheet=worksheet.title, row_number=row_number, field="加粉"
            )
            for field, fractional in zip(
                COUNT_FIELDS,
                (
                    impressions_fractional,
                    clicks_fractional,
                    uv_fractional,
                    copies_fractional,
                    adds_fractional,
                ),
                strict=True,
            ):
                fractional_counts[field] += int(fractional)

            item = aggregated.setdefault(
                keyword_key,
                {
                    "keyword_text": keyword_text,
                    "impressions": 0,
                    "clicks": 0,
                    "spend": Decimal("0"),
                    "uv": 0,
                    "copies": 0,
                    "adds": 0,
                    "source_rows": 0,
                },
            )
            item["impressions"] += impressions
            item["clicks"] += clicks
            item["spend"] += spend
            item["uv"] += uv
            item["copies"] += copies
            item["adds"] += adds
            item["source_rows"] += 1
            input_rows += 1
            current_sheet_rows += 1
        sheet_rows[worksheet.title] = current_sheet_rows

    for item in aggregated.values():
        item["spend"] = item["spend"].quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
    return aggregated, {
        "input_rows": input_rows,
        "unique_keywords": len(aggregated),
        "sheet_rows": sheet_rows,
        "fractional_values_floored": dict(fractional_counts),
    }


def import_historical_keyword_performance(
    db: Session,
    *,
    project_id: uuid.UUID,
    report_year: int,
    content: bytes,
    actor: str,
) -> dict:
    if not 2000 <= report_year <= 2100:
        raise ValueError("年份必须在2000至2100之间")
    aggregated, parse_report = parse_historical_keyword_performance(content)
    source_sha256 = hashlib.sha256(content).hexdigest()
    material_rows = db.execute(
        select(MaterialKeyword.id, MaterialKeyword.keyword_text).where(
            MaterialKeyword.project_id == project_id
        )
    ).all()
    material_by_key: dict[str, tuple[uuid.UUID, str]] = {}
    for material_id, keyword_text in material_rows:
        key = normalize_keyword_key(keyword_text)
        if key in material_by_key:
            raise ValueError(f"物料中心存在规范化后重复关键词：{keyword_text}")
        material_by_key[key] = (material_id, keyword_text)

    matched_keys = set(aggregated).intersection(material_by_key)
    unmatched_keys = set(aggregated).difference(material_by_key)
    db.execute(
        delete(MaterialKeywordPerformanceYearly).where(
            MaterialKeywordPerformanceYearly.project_id == project_id,
            MaterialKeywordPerformanceYearly.report_year == report_year,
        )
    )
    now = datetime.now(UTC)
    values = []
    totals = {
        "impressions": 0,
        "clicks": 0,
        "spend": Decimal("0"),
        "uv": 0,
        "copies": 0,
        "adds": 0,
    }
    for key in matched_keys:
        item = aggregated[key]
        material_id, _ = material_by_key[key]
        values.append(
            {
                "id": uuid.uuid4(),
                "project_id": project_id,
                "material_keyword_id": material_id,
                "report_year": report_year,
                "impressions": item["impressions"],
                "clicks": item["clicks"],
                "spend": item["spend"],
                "uv": item["uv"],
                "copies": item["copies"],
                "adds": item["adds"],
                "source_sha256": source_sha256,
                "source_rows": item["source_rows"],
                "imported_by": actor,
                "created_at": now,
                "updated_at": now,
            }
        )
        for field in totals:
            totals[field] += item[field]
    if values:
        db.execute(insert(MaterialKeywordPerformanceYearly), values)

    report = {
        **parse_report,
        "report_year": report_year,
        "source_sha256": source_sha256,
        "material_keyword_count": len(material_rows),
        "matched_keywords": len(matched_keys),
        "unmatched_keywords": len(unmatched_keys),
        "created_keywords": 0,
        "unmatched_sample": [aggregated[key]["keyword_text"] for key in sorted(unmatched_keys)[:20]],
        "matched_totals": {
            **totals,
            "spend": f"{totals['spend']:.2f}",
        },
    }
    db.add(
        AuditEvent(
            project_id=project_id,
            actor=actor,
            action="material_keyword.performance_yearly.import",
            target_type="material_keyword_performance_yearly",
            summary=f"导入{report_year}年关键词历史累计指标",
            details=report,
        )
    )
    return report
