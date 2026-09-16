from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from .models import Account
from .schemas import AccountBatchSettings


SHEET_NAME = "批量设置"
HEADERS = (
    "账户",
    "运营",
    "账户类型",
    "页面类型",
    "推广页面",
    "推广链接",
    "返点（%）",
    "钱柜账户",
)
LEGACY_SELECTOR_HEADER = "账户/账号主体"
EXAMPLE_SELECTOR_PREFIX = "示例："
FIELD_BY_HEADER = {
    "运营": "operator_name",
    "账户类型": "account_type",
    "页面类型": "page_type",
    "推广页面": "promotion_page",
    "推广链接": "promotion_link",
    "返点（%）": "rebate_rate",
    "钱柜账户": "recharge_account",
}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10_000


@dataclass
class AccountSettingsImportRow:
    row_number: int
    account_id: Any
    login_name: str
    changes: dict[str, Any]


@dataclass
class AccountSettingsImportPlan:
    file_name: str
    sha256: str
    row_count: int = 0
    matched_count: int = 0
    changed_count: int = 0
    unchanged_count: int = 0
    invalid_count: int = 0
    changed_fields: set[str] = field(default_factory=set)
    errors: list[dict[str, Any]] = field(default_factory=list)
    rows: list[AccountSettingsImportRow] = field(default_factory=list)

    def add_error(self, row: int, message: str, selector: str = "") -> None:
        self.invalid_count += 1
        self.errors.append({"row": row, "selector": selector, "message": message})

    def public(self) -> dict[str, Any]:
        return {
            "file_name": self.file_name,
            "sha256": self.sha256,
            "row_count": self.row_count,
            "matched_count": self.matched_count,
            "changed_count": self.changed_count,
            "unchanged_count": self.unchanged_count,
            "invalid_count": self.invalid_count,
            "changed_fields": sorted(self.changed_fields),
            "errors": self.errors[:200],
            "can_import": self.invalid_count == 0 and self.changed_count > 0,
        }


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _account_value(account: Account, field_name: str) -> Any:
    if field_name == "account_type":
        return account.account_type.value
    if field_name == "promotion_link":
        return account.landing_url_template
    if field_name == "lifecycle_stage":
        return account.lifecycle_override or "自动判断"
    return getattr(account, field_name)


def _values_equal(current: Any, incoming: Any) -> bool:
    if isinstance(current, Decimal) or isinstance(incoming, Decimal):
        try:
            return Decimal(str(current)) == Decimal(str(incoming))
        except (InvalidOperation, ValueError):
            return False
    current_value = current.value if hasattr(current, "value") else current
    incoming_value = incoming.value if hasattr(incoming, "value") else incoming
    return current_value == incoming_value


def _validation_message(error: ValidationError) -> str:
    messages = []
    for item in error.errors():
        field_name = str(item.get("loc", [""])[-1])
        message = str(item.get("msg", "格式错误")).removeprefix("Value error, ")
        messages.append(f"{field_name}：{message}" if field_name else message)
    return "；".join(messages)


def build_account_settings_template(accounts: Iterable[Account]) -> bytes:
    # The workbook is an input template, not an account export.  Keep the
    # argument for API compatibility but never pre-fill project accounts.
    _ = accounts
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    sheet.append((
        "示例：baidu-账户名称（请替换）",
        "王康",
        "二跳账户",
        "科普账户",
        "科普基木鱼",
        "https://example.com/page",
        12.5,
        "示例钱柜账户",
    ))

    header_fill = PatternFill("solid", fgColor="17382B")
    editable_fill = PatternFill("solid", fgColor="EAF6F0")
    warning_fill = PatternFill("solid", fgColor="FFF4CE")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    last_row = sheet.max_row
    for row in sheet.iter_rows(min_row=2, max_row=last_row):
        for cell in row:
            cell.fill = warning_fill if row[0].row == 2 else editable_fill
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{last_row}"
    widths = (34, 12, 16, 16, 18, 46, 13, 22)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

    validations = {
        "B": '"王康,王聪"',
        "C": '"一跳预埋户,一跳空户,二跳账户"',
        "D": '"科普账户,软文账户"',
        "E": '"科普基木鱼,科普全文,精华帖,中医论坛,中医秘方,快瘦汤"',
    }
    for column, formula in validations.items():
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "请从下拉选项中选择"
        validation.errorTitle = "内容不符合模板规则"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}1001")
    rebate_validation = DataValidation(
        type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True
    )
    rebate_validation.error = "返点请填写 0–100 之间的数字，例如 12.5"
    rebate_validation.showErrorMessage = True
    sheet.add_data_validation(rebate_validation)
    rebate_validation.add("G2:G1001")
    sheet.conditional_formatting.add(
        "A2:H1001", FormulaRule(formula=["LEN($A2)=0"], fill=warning_fill)
    )

    notes = workbook.create_sheet("填写说明")
    notes.append(("项目", "说明"))
    instructions = (
        ("使用方式", "“批量设置”页只保留一个示例；请替换示例账户并继续向下填写，或删除示例行后填写"),
        ("匹配规则", "账户列精确匹配当前项目的账户名或账号主体；账号主体对应多个账户时，请改用具体账户名"),
        ("空白规则", "空白单元格表示不修改，不会清空数据库中的原值"),
        ("返点", "填写百分比数字，例如 12.5 表示 12.5%，范围为 0–100"),
        ("原子导入", "导入会先完整校验；只要存在一条错误，整份文件都不会写入"),
        ("运营", "王康、王聪"),
        ("账户类型", "一跳预埋户、一跳空户、二跳账户"),
        ("页面类型", "科普账户、软文账户"),
        ("推广页面", "科普基木鱼、科普全文、精华帖、中医论坛、中医秘方、快瘦汤"),
    )
    for row in instructions:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 88
    notes.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_account_settings_workbook(
    content: bytes,
    file_name: str,
    accounts: Iterable[Account],
) -> AccountSettingsImportPlan:
    plan = AccountSettingsImportPlan(
        file_name=file_name,
        sha256=hashlib.sha256(content).hexdigest(),
    )
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as error:
        plan.add_error(0, f"无法读取工作簿：{error}")
        return plan
    if SHEET_NAME not in workbook.sheetnames:
        plan.add_error(0, f"缺少“{SHEET_NAME}”工作表，请使用系统下载的模板")
        return plan
    sheet = workbook[SHEET_NAME]
    header_positions = {_text(cell.value): cell.column for cell in sheet[1] if _text(cell.value)}
    selector_header = "账户" if "账户" in header_positions else LEGACY_SELECTOR_HEADER
    missing = [header for header in HEADERS[1:] if header not in header_positions]
    if selector_header not in header_positions:
        missing.insert(0, "账户")
    if missing:
        plan.add_error(1, f"缺少列：{'、'.join(missing)}")
        return plan

    account_rows = list(accounts)
    by_selector: dict[str, list[Account]] = {}
    for account in account_rows:
        by_selector.setdefault(account.login_name, []).append(account)
        if account.account_subject:
            by_selector.setdefault(account.account_subject, []).append(account)

    seen_account_ids: set[Any] = set()
    for row_number in range(2, sheet.max_row + 1):
        if plan.row_count >= MAX_ROWS:
            plan.add_error(row_number, f"有效数据超过 {MAX_ROWS} 行上限")
            break
        values = {
            header: sheet.cell(row_number, column).value
            for header, column in header_positions.items()
        }
        if not any(_text(value) for value in values.values()):
            continue
        selector = _text(values[selector_header])
        if selector.startswith(EXAMPLE_SELECTOR_PREFIX):
            continue
        plan.row_count += 1
        if not selector:
            plan.add_error(row_number, "账户不能为空")
            continue

        candidates = list({candidate.id: candidate for candidate in by_selector.get(selector, [])}.values())
        if not candidates:
            plan.add_error(row_number, "未在当前项目精确匹配到账户", selector)
            continue
        if len(candidates) > 1:
            plan.add_error(row_number, "账号主体匹配到多个账户，请填写具体账户名", selector)
            continue
        account = candidates[0]
        if account.id in seen_account_ids:
            plan.add_error(row_number, "该账户在工作簿中重复出现", selector)
            continue
        seen_account_ids.add(account.id)
        plan.matched_count += 1

        raw_changes: dict[str, Any] = {}
        for header, field_name in FIELD_BY_HEADER.items():
            cell = sheet.cell(row_number, header_positions[header])
            value = cell.value
            if _text(value) == "":
                continue
            if field_name == "rebate_rate":
                if isinstance(value, (int, float, Decimal)) and "%" in (cell.number_format or ""):
                    value = Decimal(str(value)) * 100
                else:
                    text_value = _text(value).removesuffix("%").strip()
                    try:
                        value = Decimal(text_value)
                    except InvalidOperation:
                        plan.add_error(row_number, "返点必须是 0–100 之间的数字", selector)
                        raw_changes = {}
                        break
            else:
                value = _text(value)
            raw_changes[field_name] = value
        if plan.errors and plan.errors[-1]["row"] == row_number:
            continue
        if not raw_changes:
            plan.unchanged_count += 1
            continue
        try:
            validated = AccountBatchSettings(selectors=[selector], **raw_changes)
        except ValidationError as error:
            plan.add_error(row_number, _validation_message(error), selector)
            continue

        changes: dict[str, Any] = {}
        for field_name in FIELD_BY_HEADER.values():
            incoming = getattr(validated, field_name)
            if incoming is None:
                continue
            if not _values_equal(_account_value(account, field_name), incoming):
                changes[field_name] = incoming
        if not changes:
            plan.unchanged_count += 1
            continue
        plan.changed_count += 1
        plan.changed_fields.update(changes)
        plan.rows.append(AccountSettingsImportRow(
            row_number=row_number,
            account_id=account.id,
            login_name=account.login_name,
            changes=changes,
        ))
    return plan


def attachment_header(file_name: str) -> str:
    return f"attachment; filename=account_batch_settings_template.xlsx; filename*=UTF-8''{quote(file_name)}"
