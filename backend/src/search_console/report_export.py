from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT_HEADERS = (
    "日期",
    "账户",
    "账户ID",
    "余额",
    "成本判断",
    "运营",
    "账户类型",
    "页面类型",
    "推广页面",
    "推广链接",
    "返点",
    "钱柜账户",
    "展现",
    "点击",
    "消费",
    "UV",
    "复制",
    "加粉",
    "CPC",
    "UV成本",
    "复制成本",
    "加粉成本",
    "现金消费",
    "现金复制成本",
    "现金加粉成本",
)


def _safe_text(value: object) -> str:
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _decimal(value: object) -> Decimal | None:
    return None if value is None else Decimal(str(value))


def build_daily_report_xlsx(rows: list[dict]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.title = "账户报表"
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24
    widths = (12, 24, 14, 12, 14, 12, 16, 14, 20, 42, 11, 20) + (12,) * 13
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header_cells = []
    for header in REPORT_HEADERS:
        cell = WriteOnlyCell(sheet, value=header)
        cell.fill = PatternFill("solid", fgColor="087A55")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cells.append(cell)
    sheet.append(header_cells)

    for row in rows:
        rebate_rate = _decimal(row.get("rebate_rate"))
        values = (
            row.get("date"),
            _safe_text(row.get("account_name")),
            row.get("account_id"),
            _decimal(row.get("balance")),
            _safe_text(row.get("cost_status")),
            _safe_text(row.get("operator_name")),
            _safe_text(row.get("account_type")),
            _safe_text(row.get("page_type")),
            _safe_text(row.get("promotion_page")),
            _safe_text(row.get("promotion_link")),
            None if rebate_rate is None else rebate_rate / Decimal("100"),
            _safe_text(row.get("recharge_account")),
            row.get("impressions"),
            row.get("clicks"),
            _decimal(row.get("spend")),
            row.get("uv"),
            row.get("copies"),
            row.get("adds"),
            _decimal(row.get("cpc")),
            _decimal(row.get("uv_cost")),
            _decimal(row.get("copy_cost")),
            _decimal(row.get("add_cost")),
            _decimal(row.get("cash_spend")),
            _decimal(row.get("cash_copy_cost")),
            _decimal(row.get("cash_add_cost")),
        )
        cells = [WriteOnlyCell(sheet, value=value) for value in values]
        cells[0].number_format = "yyyy-mm-dd"
        cells[10].number_format = "0.00%"
        for column_index in (4, 15, 19, 20, 21, 22, 23, 24, 25):
            cells[column_index - 1].number_format = '¥0.00'
        link = str(row.get("promotion_link") or "").strip()
        parsed = urlparse(link)
        if parsed.scheme in {"http", "https"} and parsed.netloc:
            cells[9].hyperlink = link
            cells[9].style = "Hyperlink"
        sheet.append(cells)

    sheet.auto_filter.ref = f"A1:Y{len(rows) + 1}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
